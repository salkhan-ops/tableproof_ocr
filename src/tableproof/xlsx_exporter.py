from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Side

from tableproof.models import TableGrid


def export_table_to_xlsx(table_grid: TableGrid, output_path: Path | str) -> None:
    output_path = Path(output_path)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table"

    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    covered_coordinates = {
        tuple(coord)
        for span in table_grid.merged_ranges()
        for coord in span["covered_coordinates"][1:]
    }

    for row in table_grid.rows:
        for cell in row.cells:
            if cell.covered_by is not None or (cell.row_index, cell.column_index) in covered_coordinates:
                continue
            excel_cell = worksheet.cell(
                row=cell.row_index + 1,
                column=cell.column_index + 1,
                value=cell.text,
            )
            excel_cell.border = border
            if cell.row_span > 1 or cell.column_span > 1:
                worksheet.merge_cells(
                    start_row=cell.row_index + 1,
                    start_column=cell.column_index + 1,
                    end_row=cell.row_index + cell.row_span,
                    end_column=cell.column_index + cell.column_span,
                )

    diagnostics = workbook.create_sheet("Diagnostics")
    diagnostics.append(["type", "row", "column", "confidence", "code", "message", "word_indices"])

    for row in table_grid.rows:
        for cell in row.cells:
            if cell.covered_by is not None or (cell.row_index, cell.column_index) in covered_coordinates:
                continue
            if cell.text and cell.confidence < 0.75:
                diagnostics.append(
                    [
                        "LOW_CONFIDENCE_CELL",
                        cell.row_index,
                        cell.column_index,
                        cell.confidence,
                        "",
                        cell.text,
                        ",".join(str(index) for index in cell.source_word_indices),
                    ]
                )

    for warning in table_grid.warnings:
        diagnostics.append(
            [
                "WARNING",
                warning.row_index,
                warning.column_index,
                "",
                warning.code,
                warning.message,
                ",".join(str(index) for index in warning.word_indices or []),
            ]
        )

    notes = workbook.create_sheet("Notes")
    notes.append(["classification", "text", "confidence", "bounding_box", "reason_codes"])
    for block in table_grid.metrics.get("excluded_text_blocks", []):
        notes.append(
            [
                block.get("classification", ""),
                block.get("text", ""),
                block.get("confidence", ""),
                str(block.get("bounding_box", "")),
                ";".join(block.get("reason_codes", [])),
            ]
        )

    raw = workbook.create_sheet("Raw Visual Lines")
    raw.append(["line_index", "row_index", "text", "bounding_box", "source_word_indices", "populated_columns"])
    for line in table_grid.metrics.get("raw_visual_lines", []):
        raw.append(
            [
                line.get("line_index"),
                line.get("row_index"),
                line.get("text"),
                str(line.get("bounding_box")),
                ",".join(str(index) for index in line.get("source_word_indices", [])),
                ",".join(str(index) for index in line.get("populated_columns", [])),
            ]
        )

    decisions = workbook.create_sheet("Continuation Decisions")
    decisions.append(
        [
            "source line",
            "target row",
            "target column",
            "candidate type",
            "selected score",
            "second-best score",
            "score margin",
            "confidence",
            "joining mode",
            "reason codes",
        ]
    )
    for decision in table_grid.metrics.get("selected_decisions", []):
        selected = decision.get("selected_candidate", {})
        alternatives = decision.get("all_considered_candidates", [])
        second_score = alternatives[1].get("score") if len(alternatives) > 1 else ""
        decisions.append(
            [
                ",".join(str(index) for index in selected.get("source_line_indices", [])),
                selected.get("target_row_index"),
                selected.get("target_column_index"),
                selected.get("candidate_type"),
                selected.get("score"),
                second_score,
                decision.get("score_margin"),
                decision.get("confidence"),
                selected.get("joining_mode"),
                ";".join(decision.get("reason_codes", [])),
            ]
        )

    unresolved = workbook.create_sheet("Unresolved Items")
    unresolved.append(["line_index", "text", "competing_interpretations", "reason_codes", "confidence"])
    for item in table_grid.metrics.get("unresolved_lines", []):
        unresolved.append(
            [
                item.get("line_index"),
                item.get("text"),
                ";".join(item.get("competing_interpretations", [])),
                ";".join(item.get("reason_codes", [])),
                item.get("confidence"),
            ]
        )

    span_decisions = workbook.create_sheet("Span Decisions")
    span_decisions.append(
        [
            "candidate ID",
            "anchor coordinate",
            "row span",
            "column span",
            "candidate type",
            "selected or rejected",
            "candidate score",
            "global solution score",
            "confidence",
            "parent",
            "children",
            "conflicts",
            "reason codes",
        ]
    )
    selected_ids = {
        decision.get("selected_candidate", {}).get("candidate_id")
        for decision in table_grid.metrics.get("selected_spans", [])
    }
    decisions_by_id = {
        decision.get("selected_candidate", {}).get("candidate_id"): decision
        for decision in table_grid.metrics.get("selected_spans", [])
    }
    for candidate in table_grid.metrics.get("span_candidates", []):
        decision = decisions_by_id.get(candidate.get("candidate_id"), {})
        span_decisions.append(
            [
                candidate.get("candidate_id"),
                f"{candidate.get('anchor_row')},{candidate.get('anchor_column')}",
                candidate.get("row_span"),
                candidate.get("column_span"),
                candidate.get("candidate_type"),
                "selected" if candidate.get("candidate_id") in selected_ids else "rejected",
                candidate.get("total_score"),
                decision.get("global_solution_score", ""),
                candidate.get("confidence"),
                decision.get("parent_span", ""),
                ";".join(decision.get("child_spans", [])),
                ";".join(candidate.get("conflicting_candidate_ids", [])),
                ";".join(decision.get("reason_codes", candidate.get("negative_evidence", []))),
            ]
        )

    hierarchy = workbook.create_sheet("Span Hierarchy")
    hierarchy.append(["hierarchy level", "parent span", "child span", "covered columns", "covered rows", "text", "confidence"])
    for key in ["header_hierarchy", "stub_hierarchy"]:
        for node in table_grid.metrics.get(key, {}).get("nodes", []):
            hierarchy.append(
                [
                    node.get("hierarchy_level"),
                    node.get("parent_candidate_id", ""),
                    node.get("candidate_id"),
                    ",".join(str(item) for item in node.get("covered_columns", [])),
                    ",".join(str(item) for item in node.get("covered_rows", [])),
                    node.get("text", ""),
                    node.get("confidence", ""),
                ]
            )

    covered = workbook.create_sheet("Covered Coordinates")
    covered.append(["row", "column", "anchor row", "anchor column", "coordinate type"])
    for coord in table_grid.metrics.get("covered_coordinate_map", []):
        covered.append(
            [
                coord.get("row_index"),
                coord.get("column_index"),
                coord.get("anchor_row"),
                coord.get("anchor_column"),
                coord.get("coordinate_type"),
            ]
        )

    pre_span = workbook.create_sheet("Pre-Span Grid")
    for row_cells in table_grid.metrics.get("pre_span_grid", []):
        pre_span.append([cell.get("text", "") for cell in row_cells])

    header_candidates = workbook.create_sheet("Header Candidates")
    header_candidates.append(
        [
            "candidate ID",
            "text",
            "type",
            "anchor row",
            "anchor column",
            "row span",
            "column span",
            "score",
            "confidence",
            "selected",
            "positive evidence",
            "negative evidence",
        ]
    )
    for candidate in table_grid.metrics.get("span_candidates", []):
        if "HEADER" not in str(candidate.get("candidate_type", "")):
            continue
        header_candidates.append(
            [
                candidate.get("candidate_id"),
                candidate.get("text"),
                candidate.get("candidate_type"),
                candidate.get("anchor_row"),
                candidate.get("anchor_column"),
                candidate.get("row_span"),
                candidate.get("column_span"),
                candidate.get("total_score"),
                candidate.get("confidence"),
                candidate.get("candidate_id") in selected_ids,
                ";".join(candidate.get("positive_evidence", [])),
                ";".join(candidate.get("negative_evidence", [])),
            ]
        )

    header_hierarchy = workbook.create_sheet("Header Hierarchy")
    header_hierarchy.append(
        [
            "node_id",
            "text",
            "logical_level",
            "visual_band",
            "start_column",
            "end_column",
            "direct_parent",
            "direct_children",
            "ancestors",
            "descendants",
            "terminal_leaf",
            "pass_through",
            "confidence",
        ]
    )
    for node in table_grid.metrics.get("header_hierarchy", {}).get("nodes", []):
        covered_columns = node.get("covered_columns", [])
        covered_rows = node.get("covered_rows", [])
        direct_children = node.get("child_candidate_ids", [])
        header_hierarchy.append(
            [
                node.get("candidate_id"),
                node.get("text"),
                node.get("hierarchy_level"),
                min(covered_rows) if covered_rows else "",
                min(covered_columns) if covered_columns else "",
                max(covered_columns) if covered_columns else "",
                node.get("parent_candidate_id", ""),
                ";".join(direct_children),
                "",
                ";".join(direct_children),
                not bool(direct_children),
                "",
                node.get("confidence", ""),
            ]
        )

    header_partitions = workbook.create_sheet("Header Partitions")
    header_partitions.append(
        [
            "parent",
            "parent interval",
            "child intervals",
            "ordinary leaf intervals",
            "covered union",
            "unexplained intervals",
            "partition complete",
            "partition confidence",
        ]
    )
    for node in table_grid.metrics.get("header_hierarchy", {}).get("nodes", []):
        children = node.get("child_candidate_ids", [])
        covered_columns = node.get("covered_columns", [])
        if not children or not covered_columns:
            continue
        header_partitions.append(
            [
                node.get("candidate_id"),
                f"{min(covered_columns)}-{max(covered_columns)}",
                ";".join(children),
                "",
                ",".join(str(column) for column in covered_columns),
                "",
                True,
                node.get("confidence", ""),
            ]
        )

    header_alternatives = workbook.create_sheet("Header Alternatives")
    header_alternatives.append(["candidate ID", "text", "type", "score", "confidence", "rejected reason", "conflicts"])
    for candidate in table_grid.metrics.get("rejected_spans", []):
        if "HEADER" not in str(candidate.get("candidate_type", "")):
            continue
        header_alternatives.append(
            [
                candidate.get("candidate_id"),
                candidate.get("text"),
                candidate.get("candidate_type"),
                candidate.get("total_score"),
                candidate.get("confidence"),
                candidate.get("rejected_reason", ""),
                ";".join(candidate.get("conflicting_candidate_ids", [])),
            ]
        )

    header_stability = workbook.create_sheet("Header Stability")
    header_stability.append(["perturbation", "selected spans", "selected hierarchy", "changed decisions", "status"])
    for row in table_grid.metrics.get("header_stability_results", table_grid.metrics.get("stability_results", [])):
        if not isinstance(row, dict):
            header_stability.append(["", "", "", "", str(row)])
            continue
        header_stability.append(
            [
                row.get("perturbation_size", ""),
                str(row.get("selected_spans", "")),
                str(row.get("selected_hierarchy", "")),
                str(row.get("changed_decisions", "")),
                row.get("stability_status", ""),
            ]
        )

    body_candidates = workbook.create_sheet("Body Group Candidates")
    body_candidates.append(["candidate ID", "anchor row", "anchor column", "row span", "type", "text", "score", "confidence", "selected", "reason codes"])
    header_rows = int(table_grid.metrics.get("estimated_header_rows") or 1)
    for candidate in table_grid.metrics.get("span_candidates", []):
        is_body_group = candidate.get("anchor_row", 0) >= header_rows and any(token in str(candidate.get("candidate_type", "")) for token in ["GROUP", "STUB"])
        if not is_body_group:
            continue
        body_candidates.append(
            [
                candidate.get("candidate_id"),
                candidate.get("anchor_row"),
                candidate.get("anchor_column"),
                candidate.get("row_span"),
                candidate.get("candidate_type"),
                candidate.get("text"),
                candidate.get("total_score"),
                candidate.get("confidence"),
                candidate.get("candidate_id") in selected_ids,
                ";".join(candidate.get("positive_evidence", []) + candidate.get("negative_evidence", [])),
            ]
        )

    body_decisions = workbook.create_sheet("Body Group Decisions")
    body_decisions.append(
        [
            "candidate ID",
            "anchor row",
            "anchor column",
            "row start",
            "row end",
            "row span",
            "candidate type",
            "selected",
            "score",
            "confidence",
            "direct parent",
            "direct children",
            "competing candidate",
            "score margin",
            "reason codes",
        ]
    )
    for decision in table_grid.metrics.get("selected_spans", []):
        candidate = decision.get("selected_candidate", {})
        if candidate.get("anchor_row", 0) < header_rows or not any(token in str(candidate.get("candidate_type", "")) for token in ["GROUP", "STUB"]):
            continue
        alternatives = decision.get("all_alternatives", [])
        body_decisions.append(
            [
                candidate.get("candidate_id"),
                candidate.get("anchor_row"),
                candidate.get("anchor_column"),
                candidate.get("anchor_row"),
                candidate.get("anchor_row", 0) + candidate.get("row_span", 1) - 1,
                candidate.get("row_span"),
                candidate.get("candidate_type"),
                True,
                candidate.get("total_score"),
                candidate.get("confidence"),
                decision.get("parent_span", ""),
                ";".join(decision.get("child_spans", [])),
                alternatives[1].get("candidate_id") if len(alternatives) > 1 else "",
                decision.get("score_margin", ""),
                ";".join(decision.get("reason_codes", [])),
            ]
        )

    stub_hierarchy_sheet = workbook.create_sheet("Stub Hierarchy")
    stub_hierarchy_sheet.append(["node ID", "text", "stub column", "indentation level", "hierarchy level", "row start", "row end", "direct parent", "direct children", "ancestors", "descendants", "confidence"])
    for node in table_grid.metrics.get("stub_hierarchy", {}).get("nodes", []):
        rows = node.get("covered_rows", [])
        columns = node.get("covered_columns", [])
        children = node.get("child_candidate_ids", [])
        stub_hierarchy_sheet.append(
            [
                node.get("candidate_id"),
                node.get("text"),
                min(columns) if columns else "",
                min(columns) if columns else "",
                node.get("hierarchy_level", ""),
                min(rows) if rows else "",
                max(rows) if rows else "",
                node.get("parent_candidate_id", ""),
                ";".join(children),
                "",
                ";".join(children),
                node.get("confidence", ""),
            ]
        )

    body_alternatives = workbook.create_sheet("Body Alternatives")
    body_alternatives.append(["candidate ID", "text", "type", "anchor row", "anchor column", "row span", "score", "confidence", "rejected reason"])
    for candidate in table_grid.metrics.get("rejected_spans", []):
        if candidate.get("anchor_row", 0) < header_rows or not any(token in str(candidate.get("candidate_type", "")) for token in ["GROUP", "STUB"]):
            continue
        body_alternatives.append(
            [
                candidate.get("candidate_id"),
                candidate.get("text"),
                candidate.get("candidate_type"),
                candidate.get("anchor_row"),
                candidate.get("anchor_column"),
                candidate.get("row_span"),
                candidate.get("total_score"),
                candidate.get("confidence"),
                candidate.get("rejected_reason", ""),
            ]
        )

    body_stability = workbook.create_sheet("Body Stability")
    body_stability.append(["perturbation", "selected row spans", "selected hierarchy", "changed decisions", "status"])
    for row in table_grid.metrics.get("body_stability_results", []):
        body_stability.append(
            [
                row.get("perturbation_size", ""),
                str(row.get("selected_spans", "")),
                str(row.get("selected_hierarchy", "")),
                str(row.get("changed_decisions", "")),
                row.get("stability_status", ""),
            ]
        )

    _add_regression_panel_sheets(workbook, table_grid, selected_ids)
    _add_mixed_span_sheets(workbook, table_grid, selected_ids)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _add_mixed_span_sheets(workbook: Workbook, table_grid: TableGrid, selected_span_ids: set[str]) -> None:
    mixed_candidates = [
        candidate
        for candidate in table_grid.metrics.get("span_candidates", []) + table_grid.metrics.get("rejected_spans", [])
        if "MIXED" in str(candidate.get("candidate_type", ""))
    ]

    candidates_sheet = workbook.create_sheet("Mixed Span Candidates")
    candidates_sheet.append(["candidate ID", "text", "anchor row", "anchor column", "row span", "column span", "score", "confidence", "selected", "covered coordinates", "reason codes"])
    for candidate in mixed_candidates:
        candidates_sheet.append(
            [
                candidate.get("candidate_id"),
                candidate.get("text"),
                candidate.get("anchor_row"),
                candidate.get("anchor_column"),
                candidate.get("row_span"),
                candidate.get("column_span"),
                candidate.get("total_score", candidate.get("score", "")),
                candidate.get("confidence"),
                candidate.get("candidate_id") in selected_span_ids,
                str(candidate.get("covered_coordinates", [])),
                ";".join(candidate.get("positive_evidence", []) + candidate.get("negative_evidence", [])),
            ]
        )

    decisions_sheet = workbook.create_sheet("Mixed Span Decisions")
    decisions_sheet.append(["candidate ID", "selected", "global score", "score margin", "second best score", "parent", "children", "reversible"])
    for decision in table_grid.metrics.get("selected_spans", []):
        candidate = decision.get("selected_candidate", {})
        if "MIXED" not in str(candidate.get("candidate_type", "")):
            continue
        decisions_sheet.append(
            [
                candidate.get("candidate_id"),
                True,
                decision.get("global_solution_score", ""),
                decision.get("score_margin", ""),
                table_grid.metrics.get("second_best_solution", {}).get("total_score", ""),
                decision.get("parent_span", ""),
                ";".join(decision.get("child_spans", [])),
                bool(decision.get("reversible_mapping")),
            ]
        )

    hierarchy_sheet = workbook.create_sheet("Mixed Span Hierarchy")
    hierarchy_sheet.append(["candidate ID", "text", "row interval", "column interval", "direct parent", "direct children", "ancestors", "descendants", "candidate type"])
    for node in table_grid.metrics.get("mixed_span_hierarchy", {}).get("nodes", []):
        hierarchy_sheet.append(
            [
                node.get("candidate_id"),
                node.get("text"),
                str(node.get("row_interval", "")),
                str(node.get("column_interval", "")),
                node.get("direct_parent", ""),
                ";".join(node.get("direct_children", [])),
                ";".join(node.get("ancestors", [])),
                ";".join(node.get("descendants", [])),
                node.get("candidate_type", ""),
            ]
        )
    if not table_grid.metrics.get("mixed_span_hierarchy", {}).get("nodes"):
        for candidate in mixed_candidates:
            hierarchy_sheet.append(
                [
                    candidate.get("candidate_id"),
                    candidate.get("text"),
                    f"{candidate.get('anchor_row')}-{candidate.get('anchor_row', 0) + candidate.get('row_span', 1) - 1}",
                    f"{candidate.get('anchor_column')}-{candidate.get('anchor_column', 0) + candidate.get('column_span', 1) - 1}",
                    "",
                    ";".join(candidate.get("child_candidate_ids", [])),
                    "",
                    ";".join(candidate.get("child_candidate_ids", [])),
                    candidate.get("candidate_type", ""),
                ]
            )

    alternatives_sheet = workbook.create_sheet("Mixed Span Alternatives")
    alternatives_sheet.append(["candidate ID", "text", "type", "score", "selected", "conflicts", "rejected reason"])
    for candidate in mixed_candidates:
        alternatives_sheet.append(
            [
                candidate.get("candidate_id"),
                candidate.get("text"),
                candidate.get("candidate_type"),
                candidate.get("total_score", ""),
                candidate.get("candidate_id") in selected_span_ids,
                ";".join(candidate.get("conflicting_candidate_ids", [])),
                candidate.get("rejected_reason", ""),
            ]
        )

    stability_sheet = workbook.create_sheet("Mixed Span Stability")
    stability_sheet.append(["perturbation", "selected mixed spans", "changed decisions", "status"])
    for row in table_grid.metrics.get("mixed_span_stability_results", table_grid.metrics.get("stability_results", [])):
        if not isinstance(row, dict):
            continue
        stability_sheet.append(
            [
                row.get("perturbation", row.get("perturbation_size", "")),
                str(row.get("selected_mixed_spans", row.get("selected_spans", ""))),
                row.get("changed_decisions", ""),
                row.get("stability_status", ""),
            ]
        )


def _add_regression_panel_sheets(workbook: Workbook, table_grid: TableGrid, selected_span_ids: set[str]) -> None:
    matrix = table_grid.matrix()
    header_rows = int(table_grid.metrics.get("estimated_header_rows") or 1)
    selected_spans = [decision.get("selected_candidate", {}) for decision in table_grid.metrics.get("selected_spans", [])]
    full_width_rows = {
        span.get("anchor_row")
        for span in selected_spans
        if span.get("column_span", 1) == table_grid.column_count and span.get("anchor_row", 0) >= header_rows
    }

    structure = workbook.create_sheet("Regression Structure")
    structure.append(["row", "row type", "text", "non_empty_cells", "source_word_indices", "reason"])
    for row in table_grid.rows:
        texts = [cell.text for cell in row.cells]
        row_text = " ".join(text for text in texts if text)
        non_empty = sum(1 for text in texts if text)
        row_type, reason = _regression_row_type(row.row_index, texts, header_rows, full_width_rows)
        structure.append(
            [
                row.row_index,
                row_type,
                row_text,
                non_empty,
                ",".join(str(index) for cell in row.cells for index in cell.source_word_indices),
                reason,
            ]
        )

    model_columns = workbook.create_sheet("Model Columns")
    model_columns.append(["column", "anchor", "boundary_left", "boundary_right", "header_text", "body_values", "evidence"])
    anchors = table_grid.metrics.get("candidate_column_anchors", [])
    boundaries = table_grid.metrics.get("candidate_boundaries", [])
    for column_index in range(table_grid.column_count):
        body_values = [
            matrix[row_index][column_index]
            for row_index in range(header_rows, table_grid.row_count)
            if matrix[row_index][column_index]
        ]
        model_columns.append(
            [
                column_index,
                str(anchors[column_index]) if column_index < len(anchors) else "",
                str(boundaries[column_index]) if column_index < len(boundaries) else "",
                str(boundaries[column_index + 1]) if column_index + 1 < len(boundaries) else "",
                " ".join(matrix[row_index][column_index] for row_index in range(min(header_rows, table_grid.row_count)) if matrix[row_index][column_index]),
                ";".join(body_values),
                "numeric_or_model_token" if any(_looks_regression_value(value) for value in body_values) else "stub_or_header",
            ]
        )

    pairs = workbook.create_sheet("Coefficient-SE Pairs")
    pairs.append(["coefficient_row", "se_row", "column", "label", "coefficient", "standard_error", "paired"])
    for row_index in range(header_rows, max(header_rows, table_grid.row_count - 1)):
        row_values = matrix[row_index]
        next_values = matrix[row_index + 1] if row_index + 1 < table_grid.row_count else []
        if not any(_looks_regression_value(value) and not _looks_standard_error(value) for value in row_values):
            continue
        if not any(_looks_standard_error(value) for value in next_values):
            continue
        label = row_values[0] if row_values else ""
        for column_index in range(1, table_grid.column_count):
            pairs.append(
                [
                    row_index,
                    row_index + 1,
                    column_index,
                    label,
                    row_values[column_index] if column_index < len(row_values) else "",
                    next_values[column_index] if column_index < len(next_values) else "",
                    bool(column_index < len(next_values) and next_values[column_index]),
                ]
            )

    stars = workbook.create_sheet("Star Attachments")
    stars.append(["row", "column", "text", "star_text", "source_word_indices", "confidence"])
    for row in table_grid.rows:
        for cell in row.cells:
            if "*" not in cell.text:
                continue
            stars.append(
                [
                    cell.row_index,
                    cell.column_index,
                    cell.text,
                    "*" * cell.text.count("*"),
                    ",".join(str(index) for index in cell.source_word_indices),
                    cell.confidence,
                ]
            )

    alternatives = workbook.create_sheet("Regression Alternatives")
    alternatives.append(["candidate ID", "text", "type", "anchor row", "anchor column", "row span", "column span", "score", "selected", "reason codes"])
    for candidate in table_grid.metrics.get("span_candidates", []) + table_grid.metrics.get("rejected_spans", []):
        candidate_type = str(candidate.get("candidate_type", ""))
        is_regression_candidate = any(token in candidate_type for token in ["HEADER", "GROUP", "STUB"]) or candidate.get("column_span", 1) > 1
        if not is_regression_candidate:
            continue
        alternatives.append(
            [
                candidate.get("candidate_id"),
                candidate.get("text"),
                candidate_type,
                candidate.get("anchor_row"),
                candidate.get("anchor_column"),
                candidate.get("row_span"),
                candidate.get("column_span"),
                candidate.get("total_score", candidate.get("score", "")),
                candidate.get("candidate_id") in selected_span_ids,
                ";".join(candidate.get("positive_evidence", []) + candidate.get("negative_evidence", [])),
            ]
        )


def _regression_row_type(row_index: int, texts: list[str], header_rows: int, full_width_rows: set[int]) -> tuple[str, str]:
    if row_index < header_rows:
        return "HEADER", "within_estimated_header_region"
    if row_index in full_width_rows:
        return "PANEL_OR_SECTION", "selected_full_width_body_span"
    if any(_looks_standard_error(text) for text in texts):
        return "STANDARD_ERROR", "parenthesized_numeric_values"
    if any(_looks_regression_value(text) and not _looks_standard_error(text) for text in texts):
        return "COEFFICIENT_OR_SUMMARY", "numeric_or_star_values"
    if sum(1 for text in texts if text) == 1:
        return "SPECIFICATION_OR_TEXT_ROW", "single_populated_cell"
    return "BODY", "default"


def _looks_standard_error(text: str) -> bool:
    value = text.strip()
    return value.startswith("(") and value.endswith(")") and any(character.isdigit() for character in value)


def _looks_regression_value(text: str) -> bool:
    value = text.strip()
    return bool(value) and any(character.isdigit() for character in value)
