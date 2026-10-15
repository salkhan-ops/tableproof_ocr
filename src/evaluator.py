from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ERROR_ROW_SPLIT = "ROW_SPLIT_ERROR"
ERROR_ROW_MERGE = "ROW_MERGE_ERROR"
ERROR_COLUMN_SPLIT = "COLUMN_SPLIT_ERROR"
ERROR_COLUMN_MERGE = "COLUMN_MERGE_ERROR"
ERROR_CELL_GROUPING = "CELL_GROUPING_ERROR"
ERROR_HEADER_ALIGNMENT = "HEADER_ALIGNMENT_ERROR"
ERROR_SPANNING_HEADER = "SPANNING_HEADER_ERROR"
ERROR_EMPTY_CELL = "EMPTY_CELL_ERROR"
ERROR_READING_ORDER = "READING_ORDER_ERROR"
ERROR_FOOTNOTE_INTRUSION = "FOOTNOTE_INTRUSION"
ERROR_SUPERSCRIPT_OR_STAR = "SUPERSCRIPT_OR_STAR_ERROR"
ERROR_TEXT_CONTENT = "TEXT_CONTENT_ERROR"
ERROR_UNASSIGNED_WORD = "UNASSIGNED_WORD"
ERROR_OTHER_STRUCTURE = "OTHER_STRUCTURE_ERROR"
ERROR_MISSED_COLUMN_SPAN = "MISSED_COLUMN_SPAN"
ERROR_FALSE_COLUMN_SPAN = "FALSE_COLUMN_SPAN"
ERROR_WRONG_COLUMN_SPAN = "WRONG_COLUMN_SPAN"
ERROR_MISSED_ROW_SPAN = "MISSED_ROW_SPAN"
ERROR_FALSE_ROW_SPAN = "FALSE_ROW_SPAN"
ERROR_WRONG_ROW_SPAN = "WRONG_ROW_SPAN"
ERROR_OVERLAPPING_SPANS = "OVERLAPPING_SPANS"
ERROR_SPAN_BOUNDARY = "SPAN_BOUNDARY_ERROR"
ERROR_TEXT_DUPLICATED_ACROSS_SPAN = "TEXT_DUPLICATED_ACROSS_SPAN"
ERROR_COVERED_CELL_TREATED_AS_EMPTY = "COVERED_CELL_TREATED_AS_EMPTY"
ERROR_AMBIGUOUS_SPAN = "AMBIGUOUS_SPAN"


UNICODE_SPACES = re.compile(r"[\s\u00a0\u1680\u2000-\u200a\u2028\u2029\u202f\u205f\u3000]+")
NUMERIC_RE = re.compile(r"^\(?[-+]?\d[\d,]*(\.\d+)?%?\)?(\*{1,3})?$")


@dataclass(frozen=True)
class EvaluationResult:
    table_id: str
    predicted_rows: int
    expected_rows: int
    predicted_columns: int
    expected_columns: int
    exact_row_count_match: bool
    exact_column_count_match: bool
    exact_cell_text_accuracy: float
    non_empty_cell_text_accuracy: float
    numeric_cell_accuracy: float
    correct_row_and_column_placement_rate: float
    empty_cell_accuracy: float
    whole_table_exact_match: bool
    cells_requiring_manual_correction: int
    percentage_cells_requiring_correction: float
    error_categories: list[str]
    cell_errors: list[dict[str, Any]]
    physical_rows: int | None = None
    physical_columns: int | None = None
    logical_cell_count: int | None = None
    expected_merged_ranges: list[dict[str, Any]] | None = None
    predicted_merged_ranges: list[dict[str, Any]] | None = None
    exact_merged_range_accuracy: float = 1.0
    column_span_detection_precision: float = 1.0
    column_span_detection_recall: float = 1.0
    row_span_detection_precision: float = 1.0
    row_span_detection_recall: float = 1.0
    span_text_accuracy: float = 1.0
    covered_coordinate_accuracy: float = 1.0
    merged_cell_placement_accuracy: float = 1.0
    percentage_spans_requiring_correction: float = 0.0
    missed_ranges: list[dict[str, Any]] | None = None
    false_ranges: list[dict[str, Any]] | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "table_id": self.table_id,
            "predicted_rows": self.predicted_rows,
            "expected_rows": self.expected_rows,
            "predicted_columns": self.predicted_columns,
            "expected_columns": self.expected_columns,
            "exact_row_count_match": self.exact_row_count_match,
            "exact_column_count_match": self.exact_column_count_match,
            "exact_cell_text_accuracy": self.exact_cell_text_accuracy,
            "non_empty_cell_text_accuracy": self.non_empty_cell_text_accuracy,
            "numeric_cell_accuracy": self.numeric_cell_accuracy,
            "correct_row_and_column_placement_rate": self.correct_row_and_column_placement_rate,
            "empty_cell_accuracy": self.empty_cell_accuracy,
            "whole_table_exact_match": self.whole_table_exact_match,
            "cells_requiring_manual_correction": self.cells_requiring_manual_correction,
            "percentage_cells_requiring_correction": self.percentage_cells_requiring_correction,
            "error_categories": ";".join(self.error_categories),
            "physical_rows": self.physical_rows if self.physical_rows is not None else self.expected_rows,
            "physical_columns": self.physical_columns if self.physical_columns is not None else self.expected_columns,
            "logical_cell_count": self.logical_cell_count,
            "exact_merged_range_accuracy": self.exact_merged_range_accuracy,
            "column_span_detection_precision": self.column_span_detection_precision,
            "column_span_detection_recall": self.column_span_detection_recall,
            "row_span_detection_precision": self.row_span_detection_precision,
            "row_span_detection_recall": self.row_span_detection_recall,
            "span_text_accuracy": self.span_text_accuracy,
            "covered_coordinate_accuracy": self.covered_coordinate_accuracy,
            "merged_cell_placement_accuracy": self.merged_cell_placement_accuracy,
            "percentage_spans_requiring_correction": self.percentage_spans_requiring_correction,
        }


@dataclass(frozen=True)
class GroundTruthTable:
    matrix: list[list[str]]
    physical_rows: int
    physical_columns: int
    logical_cells: list[dict[str, Any]]
    merged_ranges: list[dict[str, Any]]
    header_rows: int = 0
    body_rows: int = 0


def load_csv_matrix(path: Path | str) -> list[list[str]]:
    with Path(path).open(newline="", encoding="utf-8") as handle:
        return [[cell for cell in row] for row in csv.reader(handle)]


def load_ground_truth_table(path: Path | str, json_path: Path | str | None = None) -> GroundTruthTable:
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        table = _load_xlsx_ground_truth(path)
    else:
        matrix = load_csv_matrix(path)
        physical_columns = max((len(row) for row in matrix), default=0)
        logical_cells = [
            _logical_cell(row_index, column_index, cell, 1, 1)
            for row_index, row in enumerate(_normalize_matrix(matrix))
            for column_index, cell in enumerate(row)
        ]
        table = GroundTruthTable(
            matrix=_normalize_matrix(matrix),
            physical_rows=len(matrix),
            physical_columns=physical_columns,
            logical_cells=logical_cells,
            merged_ranges=[],
        )
    if json_path:
        payload = json.loads(Path(json_path).read_text(encoding="utf-8"))
        matrix = _overlay_structured_ground_truth(table.matrix, payload)
        table = GroundTruthTable(
            matrix=matrix,
            physical_rows=int(payload.get("physical_rows", table.physical_rows)),
            physical_columns=int(payload.get("physical_columns", table.physical_columns)),
            logical_cells=payload.get("logical_cells", table.logical_cells),
            merged_ranges=payload.get("merged_ranges", table.merged_ranges),
            header_rows=int(payload.get("header_rows", table.header_rows)),
            body_rows=int(payload.get("body_rows", table.body_rows)),
        )
    return table


def _overlay_structured_ground_truth(matrix: list[list[str]], payload: dict[str, Any]) -> list[list[str]]:
    physical_rows = int(payload.get("physical_rows", len(matrix)))
    physical_columns = int(payload.get("physical_columns", max((len(row) for row in matrix), default=0)))
    normalized = _normalize_matrix(matrix)
    while len(normalized) < physical_rows:
        normalized.append([])
    normalized = [row + [""] * max(0, physical_columns - len(row)) for row in normalized[:physical_rows]]
    logical_cells = payload.get("logical_cells", [])
    for cell in logical_cells:
        row_index = int(cell.get("row_index", -1))
        column_index = int(cell.get("column_index", -1))
        if row_index < 0 or column_index < 0 or row_index >= physical_rows or column_index >= physical_columns:
            continue
        normalized[row_index][column_index] = normalize_cell(cell.get("text", ""))
        for covered_row, covered_column in cell.get("covered_coordinates", []):
            if (covered_row, covered_column) == (row_index, column_index):
                continue
            if 0 <= covered_row < physical_rows and 0 <= covered_column < physical_columns:
                normalized[covered_row][covered_column] = ""
    return normalized


def table_grid_to_ground_truth(table_grid: Any) -> GroundTruthTable:
    matrix = table_grid.matrix()
    logical_cells = []
    for row in table_grid.rows:
        for cell in row.cells:
            if cell.covered_by is not None:
                continue
            logical_cells.append(
                _logical_cell(
                    cell.row_index,
                    cell.column_index,
                    cell.text,
                    cell.row_span,
                    cell.column_span,
                    cell.is_header,
                    cell.is_stub,
                )
            )
    return GroundTruthTable(
        matrix=matrix,
        physical_rows=table_grid.row_count,
        physical_columns=table_grid.column_count,
        logical_cells=logical_cells,
        merged_ranges=table_grid.merged_ranges(),
    )


def evaluate_table_grid(table_id: str, predicted_table: GroundTruthTable, ground_truth: GroundTruthTable) -> EvaluationResult:
    base = evaluate_table(table_id, predicted_table.matrix, ground_truth.matrix)
    expected_spans = [_canonical_span(span) for span in ground_truth.merged_ranges]
    predicted_spans = [_canonical_span(span) for span in predicted_table.merged_ranges]
    expected_set = {_span_key(span) for span in expected_spans}
    predicted_set = {_span_key(span) for span in predicted_spans}
    exact_matches = expected_set & predicted_set
    missed = [span for span in expected_spans if _span_key(span) not in predicted_set]
    false = [span for span in predicted_spans if _span_key(span) not in expected_set]
    categories = set(base.error_categories)

    for span in missed:
        if span["column_span"] > 1:
            categories.add(ERROR_MISSED_COLUMN_SPAN)
        if span["row_span"] > 1:
            categories.add(ERROR_MISSED_ROW_SPAN)
    for span in false:
        if span["column_span"] > 1:
            categories.add(ERROR_FALSE_COLUMN_SPAN)
        if span["row_span"] > 1:
            categories.add(ERROR_FALSE_ROW_SPAN)

    expected_col_spans = [span for span in expected_spans if span["column_span"] > 1]
    predicted_col_spans = [span for span in predicted_spans if span["column_span"] > 1]
    matched_col_spans = [
        span for span in expected_col_spans if _span_geometry_key(span) in {_span_geometry_key(p) for p in predicted_col_spans}
    ]
    expected_row_spans = [span for span in expected_spans if span["row_span"] > 1]
    predicted_row_spans = [span for span in predicted_spans if span["row_span"] > 1]
    matched_row_spans = [
        span for span in expected_row_spans if _span_geometry_key(span) in {_span_geometry_key(p) for p in predicted_row_spans}
    ]
    span_text_matches = sum(
        1
        for span in expected_spans
        for predicted in predicted_spans
        if _span_geometry_key(span) == _span_geometry_key(predicted) and span["text"] == predicted["text"]
    )
    covered_truth = {_coord for span in expected_spans for _coord in map(tuple, span["covered_coordinates"])}
    covered_pred = {_coord for span in predicted_spans for _coord in map(tuple, span["covered_coordinates"])}
    span_corrections = len(missed) + len(false)

    return EvaluationResult(
        **{
            **base.__dict__,
            "physical_rows": ground_truth.physical_rows,
            "physical_columns": ground_truth.physical_columns,
            "logical_cell_count": len(ground_truth.logical_cells),
            "expected_merged_ranges": expected_spans,
            "predicted_merged_ranges": predicted_spans,
            "exact_merged_range_accuracy": 1.0 if not expected_spans else _rate(len(exact_matches), len(expected_spans)),
            "column_span_detection_precision": 1.0 if not predicted_col_spans else _rate(len(matched_col_spans), len(predicted_col_spans)),
            "column_span_detection_recall": 1.0 if not expected_col_spans else _rate(len(matched_col_spans), len(expected_col_spans)),
            "row_span_detection_precision": 1.0 if not predicted_row_spans else _rate(len(matched_row_spans), len(predicted_row_spans)),
            "row_span_detection_recall": 1.0 if not expected_row_spans else _rate(len(matched_row_spans), len(expected_row_spans)),
            "span_text_accuracy": 1.0 if not expected_spans else _rate(span_text_matches, len(expected_spans)),
            "covered_coordinate_accuracy": 1.0 if not covered_truth else _rate(len(covered_truth & covered_pred), len(covered_truth)),
            "merged_cell_placement_accuracy": 1.0 if not expected_spans else _rate(len(exact_matches), len(expected_spans)),
            "percentage_spans_requiring_correction": 0.0 if not expected_spans else _rate(span_corrections, len(expected_spans) + len(false)),
            "error_categories": sorted(categories),
            "missed_ranges": missed,
            "false_ranges": false,
        }
    )


def normalize_cell(value: object) -> str:
    return UNICODE_SPACES.sub(" ", "" if value is None else str(value)).strip()


def evaluate_table(
    table_id: str,
    predicted: list[list[str]],
    ground_truth: list[list[str]],
) -> EvaluationResult:
    normalized_pred = _normalize_matrix(predicted)
    normalized_truth = _normalize_matrix(ground_truth)
    expected_rows = len(normalized_truth)
    predicted_rows = len(normalized_pred)
    expected_columns = max((len(row) for row in normalized_truth), default=0)
    predicted_columns = max((len(row) for row in normalized_pred), default=0)
    compare_rows = max(expected_rows, predicted_rows)
    compare_columns = max(expected_columns, predicted_columns)

    expected_total = expected_rows * expected_columns
    compare_total = compare_rows * compare_columns
    exact_matches = 0
    placement_matches = 0
    non_empty_total = 0
    non_empty_matches = 0
    numeric_total = 0
    numeric_matches = 0
    empty_total = 0
    empty_matches = 0
    cell_errors: list[dict[str, Any]] = []
    categories: set[str] = set()

    if predicted_rows > expected_rows:
        categories.add(ERROR_ROW_SPLIT)
    elif predicted_rows < expected_rows:
        categories.add(ERROR_ROW_MERGE)
    if predicted_columns > expected_columns:
        categories.add(ERROR_COLUMN_SPLIT)
    elif predicted_columns < expected_columns:
        categories.add(ERROR_COLUMN_MERGE)

    for row_index in range(compare_rows):
        for column_index in range(compare_columns):
            truth = _matrix_get(normalized_truth, row_index, column_index)
            pred = _matrix_get(normalized_pred, row_index, column_index)
            if row_index < expected_rows and column_index < expected_columns:
                if truth == "":
                    empty_total += 1
                    if pred == "":
                        empty_matches += 1
                else:
                    non_empty_total += 1
                    if truth == pred:
                        non_empty_matches += 1
                    if is_numeric_like(truth):
                        numeric_total += 1
                        if truth == pred:
                            numeric_matches += 1

            if truth == pred:
                if row_index < expected_rows and column_index < expected_columns:
                    exact_matches += 1
                placement_matches += 1
                continue

            category = classify_cell_error(truth, pred, row_index, column_index, expected_rows, expected_columns)
            categories.add(category)
            cell_errors.append(
                {
                    "table_id": table_id,
                    "row_index": row_index,
                    "column_index": column_index,
                    "expected": truth,
                    "predicted": pred,
                    "error_type": category,
                }
            )

    whole_match = (
        predicted_rows == expected_rows
        and predicted_columns == expected_columns
        and len(cell_errors) == 0
    )
    corrections = len(cell_errors)
    correction_denominator = max(compare_total, 1)

    return EvaluationResult(
        table_id=table_id,
        predicted_rows=predicted_rows,
        expected_rows=expected_rows,
        predicted_columns=predicted_columns,
        expected_columns=expected_columns,
        exact_row_count_match=predicted_rows == expected_rows,
        exact_column_count_match=predicted_columns == expected_columns,
        exact_cell_text_accuracy=_rate(exact_matches, max(expected_total, 1)),
        non_empty_cell_text_accuracy=_rate(non_empty_matches, max(non_empty_total, 1)),
        numeric_cell_accuracy=1.0 if numeric_total == 0 else _rate(numeric_matches, numeric_total),
        correct_row_and_column_placement_rate=_rate(placement_matches, max(compare_total, 1)),
        empty_cell_accuracy=1.0 if empty_total == 0 else _rate(empty_matches, empty_total),
        whole_table_exact_match=whole_match,
        cells_requiring_manual_correction=corrections,
        percentage_cells_requiring_correction=_rate(corrections, correction_denominator),
        error_categories=sorted(categories),
        cell_errors=cell_errors,
    )


def is_numeric_like(value: str) -> bool:
    return bool(NUMERIC_RE.fullmatch(normalize_cell(value)))


def classify_cell_error(
    truth: str,
    pred: str,
    row_index: int,
    column_index: int,
    expected_rows: int,
    expected_columns: int,
) -> str:
    if "note" in pred.lower() and "note" not in truth.lower():
        return ERROR_FOOTNOTE_INTRUSION
    if row_index >= expected_rows:
        return ERROR_ROW_SPLIT
    if column_index >= expected_columns:
        return ERROR_COLUMN_SPLIT
    if truth == "" or pred == "":
        return ERROR_EMPTY_CELL
    if "*" in truth or "*" in pred:
        return ERROR_SUPERSCRIPT_OR_STAR
    if ("(" in truth or ")" in truth or "(" in pred or ")" in pred) and truth != pred:
        return ERROR_TEXT_CONTENT
    return ERROR_TEXT_CONTENT


def table_passes(result: EvaluationResult) -> bool:
    base_pass = (
        result.exact_row_count_match
        and result.exact_column_count_match
        and result.numeric_cell_accuracy >= 0.99
        and result.correct_row_and_column_placement_rate >= 0.98
        and result.percentage_cells_requiring_correction <= 0.05
    )
    if result.expected_merged_ranges:
        return (
            base_pass
            and result.exact_merged_range_accuracy == 1.0
            and result.column_span_detection_precision == 1.0
            and result.column_span_detection_recall == 1.0
            and result.row_span_detection_precision == 1.0
            and result.row_span_detection_recall == 1.0
            and result.span_text_accuracy == 1.0
            and result.percentage_spans_requiring_correction <= 0.05
        )
    return base_pass


def _load_xlsx_ground_truth(path: Path) -> GroundTruthTable:
    workbook = load_workbook(path)
    worksheet = workbook.active
    physical_rows = worksheet.max_row
    physical_columns = worksheet.max_column
    matrix = [["" for _ in range(physical_columns)] for _ in range(physical_rows)]
    merged_lookup: dict[tuple[int, int], dict[str, Any]] = {}
    covered = set()
    for merged_range in worksheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        row_index = min_row - 1
        column_index = min_col - 1
        row_span = max_row - min_row + 1
        column_span = max_col - min_col + 1
        text = normalize_cell(worksheet.cell(min_row, min_col).value)
        span = _logical_cell(row_index, column_index, text, row_span, column_span, row_index < 2, column_index == 0)
        merged_lookup[(row_index, column_index)] = span
        covered.update(tuple(coord) for coord in span["covered_coordinates"][1:])
    logical_cells = []
    for row_index in range(physical_rows):
        for column_index in range(physical_columns):
            if (row_index, column_index) in covered:
                matrix[row_index][column_index] = ""
                continue
            value = normalize_cell(worksheet.cell(row_index + 1, column_index + 1).value)
            matrix[row_index][column_index] = value
            logical_cells.append(
                merged_lookup.get(
                    (row_index, column_index),
                    _logical_cell(row_index, column_index, value, 1, 1, row_index < 2, column_index == 0),
                )
            )
    return GroundTruthTable(
        matrix=matrix,
        physical_rows=physical_rows,
        physical_columns=physical_columns,
        logical_cells=logical_cells,
        merged_ranges=list(merged_lookup.values()),
        header_rows=2,
        body_rows=max(0, physical_rows - 2),
    )


def _logical_cell(
    row_index: int,
    column_index: int,
    text: str,
    row_span: int,
    column_span: int,
    is_header: bool = False,
    is_stub: bool = False,
) -> dict[str, Any]:
    return {
        "row_index": row_index,
        "column_index": column_index,
        "text": normalize_cell(text),
        "row_span": row_span,
        "column_span": column_span,
        "covered_coordinates": [
            [r, c]
            for r in range(row_index, row_index + row_span)
            for c in range(column_index, column_index + column_span)
        ],
        "is_header": is_header,
        "is_stub": is_stub,
        "bounding_box": None,
    }


def _canonical_span(span: dict[str, Any]) -> dict[str, Any]:
    return {
        "row_index": int(span["row_index"]),
        "column_index": int(span["column_index"]),
        "text": normalize_cell(span.get("text", "")),
        "row_span": int(span.get("row_span", 1)),
        "column_span": int(span.get("column_span", 1)),
        "covered_coordinates": [list(coord) for coord in span.get("covered_coordinates", [])],
        "is_header": bool(span.get("is_header", False)),
        "is_stub": bool(span.get("is_stub", False)),
    }


def _span_geometry_key(span: dict[str, Any]) -> tuple[Any, ...]:
    return (
        span["row_index"],
        span["column_index"],
        span["row_span"],
        span["column_span"],
        tuple(tuple(coord) for coord in span["covered_coordinates"]),
    )


def _span_key(span: dict[str, Any]) -> tuple[Any, ...]:
    return _span_geometry_key(span) + (span["text"],)


def _normalize_matrix(matrix: list[list[str]]) -> list[list[str]]:
    width = max((len(row) for row in matrix), default=0)
    return [[normalize_cell(cell) for cell in row] + [""] * (width - len(row)) for row in matrix]


def _matrix_get(matrix: list[list[str]], row_index: int, column_index: int) -> str:
    if row_index >= len(matrix) or column_index >= len(matrix[row_index]):
        return ""
    return matrix[row_index][column_index]


def _rate(numerator: int, denominator: int) -> float:
    return round(numerator / denominator, 6)
