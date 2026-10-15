from __future__ import annotations

from pathlib import Path

import json
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import load_workbook

from models import BoundingBox, WordBox
from table_reconstructor import (
    WARNING_AMBIGUOUS_COLUMN,
    WARNING_EXPECTED_EMPTY_CELL_PATTERN,
    WARNING_FOOTNOTE_EXCLUDED,
    WARNING_IRREGULAR_COLUMN_COUNT,
    WARNING_POSSIBLE_MERGED_CELL,
    reconstruct_table_grid,
    reconstruct_from_words_file,
)

from benchmark_runner import run_benchmark


def wb(index: int, text: str, x0: float, y0: float, x1: float, y1: float) -> WordBox:
    return WordBox(index=index, text=text, bounding_box=BoundingBox(x0, y0, x1, y1))


def make_words(rows: list[list[tuple[str, float]]], y_start: float = 10.0) -> list[WordBox]:
    words: list[WordBox] = []
    index = 0
    for row_index, row in enumerate(rows):
        y0 = y_start + row_index * 22
        for text, x0 in row:
            width = max(8.0, len(text) * 5.0)
            words.append(wb(index, text, x0, y0, x0 + width, y0 + 10))
            index += 1
    return words


def codes(grid) -> set[str]:
    return {warning.code for warning in grid.warnings}


def test_simple_three_column_table() -> None:
    words = make_words(
        [
            [("Name", 10), ("Q1", 100), ("Q2", 160)],
            [("Alpha", 10), ("10", 100), ("20", 160)],
            [("Beta", 10), ("30", 100), ("40", 160)],
        ]
    )

    grid = reconstruct_table_grid(words)

    assert grid.row_count == 3
    assert grid.column_count == 3
    assert grid.matrix() == [["Name", "Q1", "Q2"], ["Alpha", "10", "20"], ["Beta", "30", "40"]]
    assert grid.to_dict() == reconstruct_table_grid(words).to_dict()


def test_exact_failing_sample_reconstructs_three_columns() -> None:
    words = [
        wb(0, "Revenue", 72.0, 83.1, 120.02, 99.58),
        wb(1, "Q1", 123.36, 83.1, 139.36, 99.58),
        wb(2, "Q2", 142.70, 83.1, 158.71, 99.58),
        wb(3, "Product", 72.0, 109.1, 113.35, 125.58),
        wb(4, "A", 116.68, 109.1, 124.69, 125.58),
        wb(5, "100", 128.02, 109.1, 148.04, 125.58),
        wb(6, "120", 151.37, 109.1, 171.39, 125.58),
        wb(7, "Product", 72.0, 135.1, 113.35, 151.58),
        wb(8, "B", 116.68, 135.1, 124.69, 151.58),
        wb(9, "90", 128.02, 135.1, 141.37, 151.58),
        wb(10, "110", 144.70, 135.1, 164.72, 151.58),
    ]

    grid = reconstruct_table_grid(words)

    assert grid.row_count == 3
    assert grid.column_count == 3
    assert grid.matrix() == [
        ["Revenue", "Q1", "Q2"],
        ["Product A", "100", "120"],
        ["Product B", "90", "110"],
    ]
    assert WARNING_AMBIGUOUS_COLUMN not in codes(grid)


def test_borderless_table() -> None:
    words = make_words(
        [
            [("Region", 12), ("Jan", 96), ("Feb", 157)],
            [("East", 11), ("5", 98), ("7", 158)],
        ]
    )

    grid = reconstruct_table_grid(words)

    assert grid.row_count == 2
    assert grid.column_count == 3
    assert grid.matrix()[1] == ["East", "5", "7"]


def test_multi_word_cells_are_grouped() -> None:
    words = [
        wb(0, "City", 10, 10, 30, 20),
        wb(1, "Value", 100, 10, 125, 20),
        wb(2, "New", 10, 32, 28, 42),
        wb(3, "York", 31, 32, 52, 42),
        wb(4, "100", 100, 32, 118, 42),
        wb(5, "Los", 10, 54, 26, 64),
        wb(6, "Angeles", 29, 54, 64, 64),
        wb(7, "200", 100, 54, 118, 64),
    ]

    grid = reconstruct_table_grid(words)

    assert grid.row_count == 3
    assert grid.column_count == 2
    assert grid.matrix()[1] == ["New York", "100"]
    assert grid.matrix()[2] == ["Los Angeles", "200"]


def test_multi_word_first_column_with_two_right_aligned_numeric_columns() -> None:
    words = [
        wb(0, "Line", 10, 10, 28, 20),
        wb(1, "A", 31, 10, 39, 20),
        wb(2, "95", 116, 10, 130, 20),
        wb(3, "1025", 176, 10, 205, 20),
        wb(4, "Line", 10, 32, 28, 42),
        wb(5, "B", 31, 32, 39, 42),
        wb(6, "105", 109, 32, 130, 42),
        wb(7, "88", 191, 32, 205, 42),
    ]

    grid = reconstruct_table_grid(words)

    assert grid.column_count == 3
    assert grid.matrix() == [["Line A", "95", "1025"], ["Line B", "105", "88"]]


def test_different_digit_widths_align_as_one_right_aligned_column() -> None:
    words = [
        wb(0, "Name", 10, 10, 31, 20),
        wb(1, "Value", 120, 10, 145, 20),
        wb(2, "A", 10, 32, 18, 42),
        wb(3, "9", 138, 32, 145, 42),
        wb(4, "B", 10, 54, 18, 64),
        wb(5, "100", 124, 54, 145, 64),
        wb(6, "C", 10, 76, 18, 86),
        wb(7, "1250", 116, 76, 145, 86),
    ]

    grid = reconstruct_table_grid(words)

    assert grid.column_count == 2
    assert [row[1] for row in grid.matrix()] == ["Value", "9", "100", "1250"]
    assert grid.metrics["anchor_alignment_types"][1] == "right"


def test_centered_headers_above_right_aligned_numeric_columns() -> None:
    words = [
        wb(0, "Item", 10, 10, 30, 20),
        wb(1, "Q1", 114, 10, 126, 20),
        wb(2, "Q2", 174, 10, 186, 20),
        wb(3, "Alpha", 10, 32, 36, 42),
        wb(4, "100", 109, 32, 130, 42),
        wb(5, "120", 169, 32, 190, 42),
        wb(6, "Beta", 10, 54, 32, 64),
        wb(7, "90", 116, 54, 130, 64),
        wb(8, "110", 169, 54, 190, 64),
    ]

    grid = reconstruct_table_grid(words)

    assert grid.column_count == 3
    assert grid.matrix()[0] == ["Item", "Q1", "Q2"]


def test_genuine_one_column_text_list_stays_one_column() -> None:
    words = [
        wb(0, "First", 10, 10, 34, 20),
        wb(1, "item", 37, 10, 57, 20),
        wb(2, "Second", 10, 32, 42, 42),
        wb(3, "item", 45, 32, 65, 42),
        wb(4, "Third", 10, 54, 36, 64),
        wb(5, "item", 39, 54, 59, 64),
    ]

    grid = reconstruct_table_grid(words)

    assert grid.column_count == 1
    assert grid.matrix() == [["First item"], ["Second item"], ["Third item"]]
    assert WARNING_AMBIGUOUS_COLUMN not in codes(grid)


def test_missing_cell_is_preserved() -> None:
    words = make_words(
        [
            [("Item", 10), ("Q1", 100), ("Q2", 160)],
            [("Alpha", 10), ("10", 100), ("20", 160)],
            [("Beta", 10), ("40", 160)],
        ]
    )

    grid = reconstruct_table_grid(words)

    assert grid.column_count == 3
    assert grid.matrix()[2] == ["Beta", "", "40"]
    assert WARNING_IRREGULAR_COLUMN_COUNT not in codes(grid)
    assert any(
        item["reason"] == WARNING_EXPECTED_EMPTY_CELL_PATTERN
        for item in grid.metrics["warning_suppression_reasons"]
    )


def test_two_level_header_flags_possible_spanning_cell() -> None:
    words = [
        wb(0, "Product", 10, 10, 46, 20),
        wb(1, "Sales", 100, 10, 190, 20),
        wb(2, "Name", 10, 32, 32, 42),
        wb(3, "Q1", 100, 32, 112, 42),
        wb(4, "Q2", 160, 32, 172, 42),
        wb(5, "Alpha", 10, 54, 36, 64),
        wb(6, "10", 100, 54, 112, 64),
        wb(7, "20", 160, 54, 172, 64),
    ]

    grid = reconstruct_table_grid(words)
    sales_cell = grid.rows[0].cells[1]

    assert grid.column_count == 3
    assert sales_cell.text == "Sales"
    assert sales_cell.column_span == 2
    assert WARNING_POSSIBLE_MERGED_CELL in codes(grid)


def test_slightly_misaligned_words_share_columns() -> None:
    words = make_words(
        [
            [("Name", 10), ("Q1", 100), ("Q2", 160)],
            [("Alpha", 13), ("10", 103), ("20", 158)],
            [("Beta", 8), ("30", 98), ("40", 162)],
        ]
    )

    grid = reconstruct_table_grid(words)

    assert grid.column_count == 3
    assert grid.matrix()[1] == ["Alpha", "10", "20"]
    assert WARNING_AMBIGUOUS_COLUMN not in codes(grid)


def test_close_column_spacing_uses_cross_row_alignment() -> None:
    words = [
        wb(0, "Metric", 10, 10, 40, 20),
        wb(1, "A", 70, 10, 78, 20),
        wb(2, "B", 91, 10, 99, 20),
        wb(3, "Foo", 10, 32, 28, 42),
        wb(4, "10", 68, 32, 82, 42),
        wb(5, "20", 91, 32, 105, 42),
        wb(6, "Bar", 10, 54, 28, 64),
        wb(7, "11", 68, 54, 82, 64),
        wb(8, "21", 91, 54, 105, 64),
    ]

    grid = reconstruct_table_grid(words)

    assert grid.column_count == 3
    assert grid.matrix()[1] == ["Foo", "10", "20"]


def test_reconstruction_is_deterministic_across_repeated_runs() -> None:
    words = make_words(
        [
            [("Name", 10), ("Q1", 100), ("Q2", 160)],
            [("Alpha", 10), ("10", 100), ("20", 160)],
            [("Beta", 10), ("30", 100), ("40", 160)],
        ]
    )

    first = reconstruct_table_grid(words).to_dict()
    second = reconstruct_table_grid(list(reversed(words))).to_dict()

    assert first == second


def test_ambiguous_column_assignment_is_warned() -> None:
    words = make_words(
        [
            [("A", 10), ("B", 100), ("C", 160)],
            [("x", 10), ("middle", 132), ("z", 160)],
        ]
    )

    grid = reconstruct_table_grid(words, column_tolerance=4)

    assert WARNING_AMBIGUOUS_COLUMN in codes(grid)


def note_words(prefix: str = "Note:", y_gap: float = 22, note_height: float = 10) -> list[WordBox]:
    words = make_words(
        [
            [("Measure", 10), ("A", 100), ("B", 160), ("Diff", 220)],
            [("Rate", 10), ("91%", 100), ("88%", 160), ("3%", 220)],
            [("Days", 10), ("14", 100), ("17", 160), ("-3", 220)],
            [("Score", 10), ("4.2", 100), ("4.0", 160), ("0.2", 220)],
        ]
    )
    y0 = 10 + 3 * 22 + y_gap
    start = len(words)
    for offset, (text, x0, width) in enumerate(
        [
            (prefix, 10, 28),
            ("Values", 42, 34),
            ("are", 80, 18),
            ("descriptive", 102, 55),
            ("and", 161, 18),
            ("unweighted.", 183, 58),
        ]
    ):
        words.append(wb(start + offset, text, x0, y0, x0 + width, y0 + note_height))
    return words


def test_borderless_table_followed_by_note_is_excluded() -> None:
    grid = reconstruct_table_grid(note_words("Note:"))

    assert grid.row_count == 4
    assert grid.column_count == 4
    assert grid.metrics["excluded_text_blocks"][0]["classification"] == "GENERAL_NOTE"
    assert WARNING_FOOTNOTE_EXCLUDED in codes(grid)


def test_borderless_table_followed_by_source_is_excluded() -> None:
    grid = reconstruct_table_grid(note_words("Source:"))

    assert grid.row_count == 4
    assert grid.metrics["excluded_text_blocks"][0]["classification"] == "SOURCE_NOTE"


def test_regression_significance_explanation_is_excluded() -> None:
    words = make_words(
        [
            [("Variable", 10), ("Model", 100), ("Model", 170)],
            [("Treatment", 10), ("0.1**", 100), ("0.2*", 170)],
            [("", 10), ("(0.03)", 100), ("(0.04)", 170)],
        ]
    )
    words.extend(
        [
            wb(9, "Significance", 10, 86, 68, 96),
            wb(10, "levels:", 72, 86, 105, 96),
            wb(11, "*", 109, 86, 114, 96),
            wb(12, "p<0.10;", 118, 86, 154, 96),
            wb(13, "**", 158, 86, 168, 96),
            wb(14, "p<0.05.", 172, 86, 210, 96),
        ]
    )

    grid = reconstruct_table_grid(words)

    assert grid.row_count == 3
    assert grid.metrics["excluded_text_blocks"][0]["classification"] == "FOOTNOTE"


def test_footnote_smaller_font_is_excluded() -> None:
    grid = reconstruct_table_grid(note_words("Note:", note_height=7))

    assert grid.row_count == 4
    assert grid.metrics["excluded_text_blocks"]


def test_footnote_same_font_size_is_excluded() -> None:
    grid = reconstruct_table_grid(note_words("Note:", note_height=10))

    assert grid.row_count == 4
    assert grid.metrics["excluded_text_blocks"]


def test_long_prose_line_crossing_columns_is_excluded() -> None:
    words = make_words(
        [
            [("Item", 10), ("A", 100), ("B", 160)],
            [("Alpha", 10), ("10", 100), ("20", 160)],
            [("Beta", 10), ("30", 100), ("40", 160)],
        ]
    )
    start = len(words)
    for offset, text in enumerate(["Figures", "in", "this", "table", "are", "unaudited."]):
        words.append(wb(start + offset, text, 10 + offset * 35, 92, 30 + offset * 35, 102))

    grid = reconstruct_table_grid(words)

    assert grid.row_count == 3
    assert grid.metrics["excluded_text_blocks"][0]["classification"] in {"GENERAL_NOTE", "UNCERTAIN_TRAILING_TEXT"}


def test_wrapped_first_column_cell_remains_inside_table() -> None:
    words = [
        wb(0, "Description", 10, 10, 62, 20),
        wb(1, "Value", 130, 10, 155, 20),
        wb(2, "Long", 10, 32, 32, 42),
        wb(3, "product", 35, 32, 68, 42),
        wb(4, "100", 130, 32, 148, 42),
        wb(5, "continued", 10, 45, 55, 55),
        wb(6, "Short", 10, 67, 36, 77),
        wb(7, "200", 130, 67, 148, 77),
    ]

    grid = reconstruct_table_grid(words)

    assert grid.row_count == 3
    assert not grid.metrics["excluded_text_blocks"]
    assert grid.matrix()[1] == ["Long product continued", "100"]
    assert grid.metrics["selected_decisions"]


def test_genuine_final_all_text_row_is_not_removed() -> None:
    words = make_words(
        [
            [("Type", 10), ("Status", 100), ("Owner", 170)],
            [("Alpha", 10), ("Open", 100), ("Team", 170)],
            [("Summary", 10), ("Complete", 100), ("Internal", 170)],
        ]
    )

    grid = reconstruct_table_grid(words)

    assert grid.row_count == 3
    assert not grid.metrics["excluded_text_blocks"]


def test_large_vertical_gap_before_note() -> None:
    grid = reconstruct_table_grid(note_words("Note:", y_gap=42))

    assert grid.row_count == 4
    assert "VERTICAL_GAP" in grid.metrics["excluded_text_blocks"][0]["reason_codes"]


def test_small_vertical_gap_before_note() -> None:
    grid = reconstruct_table_grid(note_words("Note:", y_gap=14))

    assert grid.row_count == 4
    assert grid.metrics["excluded_text_blocks"]


def test_correlation_matrix_triangular_empty_cells_suppress_irregular_warning() -> None:
    words = make_words(
        [
            [("Sales", 100), ("Costs", 160), ("Profit", 220)],
            [("Sales", 10), ("1.00", 100)],
            [("Costs", 10), ("0.82", 100), ("1.00", 160)],
            [("Profit", 10), ("0.41", 100), ("0.12", 160), ("1.00", 220)],
        ]
    )

    grid = reconstruct_table_grid(words)

    assert WARNING_IRREGULAR_COLUMN_COUNT not in codes(grid)
    assert any(
        item["reason"] == "TRIANGULAR_MATRIX_PATTERN"
        for item in grid.metrics["warning_suppression_reasons"]
    )


def test_regression_intentional_blank_cells_suppress_irregular_warning() -> None:
    words = [
        wb(0, "Variable", 10, 10, 50, 20),
        wb(1, "Model", 100, 10, 130, 20),
        wb(2, "Treatment", 10, 32, 55, 42),
        wb(3, "0.12**", 100, 32, 132, 42),
        wb(4, "(0.04)", 100, 54, 132, 64),
    ]

    grid = reconstruct_table_grid(words)

    assert WARNING_IRREGULAR_COLUMN_COUNT not in codes(grid)
    assert grid.matrix()[2] == ["", "(0.04)"]


def test_multi_level_column_spanning_header_range_is_preserved() -> None:
    words = [
        wb(0, "Category", 10, 10, 50, 20),
        wb(1, "Q1 Results", 100, 10, 150, 20),
        wb(2, "Revenue", 100, 32, 138, 42),
        wb(3, "Cost", 160, 32, 182, 42),
        wb(4, "A", 10, 54, 18, 64),
        wb(5, "10", 100, 54, 112, 64),
        wb(6, "7", 160, 54, 168, 64),
    ]

    grid = reconstruct_table_grid(words)

    assert any(span["text"] == "Q1 Results" and span["column_span"] == 2 for span in grid.merged_ranges())


def test_vertical_row_spanning_stub_label_range_is_preserved() -> None:
    words = [
        wb(0, "Category", 10, 10, 50, 20),
        wb(1, "Value", 100, 10, 126, 20),
        wb(2, "Group", 10, 32, 38, 42),
        wb(3, "A", 41, 32, 49, 42),
        wb(4, "10", 100, 32, 112, 42),
        wb(5, "20", 100, 54, 112, 64),
        wb(6, "30", 100, 76, 112, 86),
    ]

    grid = reconstruct_table_grid(words)

    assert any(span["text"] == "Group A" and span["row_span"] == 3 for span in grid.merged_ranges())


def test_footnote_text_is_preserved_in_notes_worksheet(tmp_path: Path) -> None:
    words = [
        {"text": word.text, "bbox": word.bounding_box.to_dict()}
        for word in note_words("Note:")
    ]
    words_path = tmp_path / "words.json"
    words_path.write_text(json.dumps(words), encoding="utf-8")

    result = reconstruct_from_words_file(words_path, output_dir=tmp_path)
    workbook = load_workbook(result.xlsx_path)
    notes = workbook["Notes"]

    assert notes.cell(2, 1).value == "GENERAL_NOTE"
    assert "Note: Values are descriptive" in notes.cell(2, 2).value


def test_boundary_decisions_are_deterministic() -> None:
    words = note_words("Source:")

    first = reconstruct_table_grid(words).metrics["table_region"]
    second = reconstruct_table_grid(list(reversed(words))).metrics["table_region"]

    assert first == second


def test_manifest_benchmark_tables_remain_passing() -> None:
    result = run_benchmark("data/benchmark_manifest.csv")

    assert result["summary"]["tables_passed"] == 5
    assert result["summary"]["whole_table_exact_matches"] == 5
