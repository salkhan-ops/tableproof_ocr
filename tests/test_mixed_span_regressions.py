from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))

from tableproof.models import BoundingBox, Cell, TableRow
from tableproof.span_candidate_generator import generate_span_candidates
from tableproof.span_constraint_solver import solve_span_candidates
from tableproof.table_reconstructor import load_words_json, reconstruct_from_words_file, reconstruct_table_grid


ROOT = Path(__file__).resolve().parents[1]


def _cell(row: int, column: int, text: str) -> Cell:
    box = BoundingBox(50 + column * 70, 20 + row * 15, 90 + column * 70, 30 + row * 15) if text else None
    return Cell(row, column, text, box, [row * 10 + column] if text else [])


def _rows(values: list[list[str]]) -> list[TableRow]:
    return [
        TableRow(row_index, [_cell(row_index, column, text) for column, text in enumerate(row)], None, 1.0)
        for row_index, row in enumerate(values)
    ]


def _span_key(span: dict) -> tuple:
    return (
        span["row_index"],
        span["column_index"],
        span["row_span"],
        span["column_span"],
        span["text"],
        tuple(tuple(coord) for coord in span["covered_coordinates"]),
    )


def test_simple_two_by_two_mixed_span_candidate_is_selected() -> None:
    candidates = generate_span_candidates(
        _rows([
            ["Descriptor", "", "Results"],
            ["", "", "Q1"],
            ["Product", "Line", "100"],
        ]),
        {"estimated_header_rows": 2},
    )

    mixed = [candidate for candidate in candidates if candidate.candidate_type == "MIXED_ROW_COLUMN_SPAN"]
    selected, _rejected, _conflicts, _solution = solve_span_candidates(candidates)

    assert [(candidate.anchor_row, candidate.anchor_column, candidate.row_span, candidate.column_span) for candidate in mixed] == [(0, 0, 2, 2)]
    assert any(candidate.candidate_type == "MIXED_ROW_COLUMN_SPAN" for candidate in selected)


def test_populated_coordinate_rejects_false_mixed_span() -> None:
    candidates = generate_span_candidates(
        _rows([
            ["Descriptor", "Other", "Results"],
            ["", "", "Q1"],
            ["Product", "Line", "100"],
        ]),
        {"estimated_header_rows": 2},
    )

    mixed = [candidate for candidate in candidates if candidate.candidate_type == "MIXED_ROW_COLUMN_SPAN"]

    assert mixed == []


def test_mixed_span_beats_conflicting_single_direction_stub_span() -> None:
    candidates = generate_span_candidates(
        _rows([
            ["Descriptor", "", "Results", ""],
            ["", "", "Q1", ""],
            ["", "", "Revenue", "Cost"],
            ["Consumer", "Product A", "100", "70"],
        ]),
        {"estimated_header_rows": 3},
    )
    selected, rejected, _conflicts, solution = solve_span_candidates(candidates)

    assert "mixed_row_column_span_0_0_3x2" in solution.selected_candidate_ids
    assert any("stub_header_span_0_0_3x1" in candidate.conflicting_candidate_ids for candidate in selected)
    assert any(candidate.candidate_id == "stub_header_span_0_0_3x1" for candidate in rejected)
    assert all(candidate.candidate_id != "stub_header_span_0_0_3x1" for candidate in selected)


def test_span_06_exact_mixed_row_column_reconstruction() -> None:
    expected = json.loads((ROOT / "data/ground_truth/span_06_ground_truth.json").read_text(encoding="utf-8"))
    words = load_words_json(ROOT / "outputs/extracted_tables/span_06_page_001_words.json")

    grid = reconstruct_table_grid(words, debug=True)

    assert grid.row_count == expected["physical_rows"]
    assert grid.column_count == expected["physical_columns"]
    assert {_span_key(span) for span in grid.merged_ranges()} == {
        _span_key(span) for span in expected["merged_ranges"]
    }
    assert grid.matrix()[3] == ["Consumer", "Product A", "100", "70", "120", "80"]
    assert grid.metrics["no_word_loss"]["passed"]


def test_span_06_small_coordinate_perturbation_is_stable() -> None:
    words = load_words_json(ROOT / "outputs/extracted_tables/span_06_page_001_words.json")
    reference = reconstruct_table_grid(words, debug=True)
    reference_signature = (reference.matrix(), sorted(_span_key(span) for span in reference.merged_ranges()))
    perturbed = []
    for word in words:
        delta = 0.5 if word.index % 2 == 0 else -0.5
        box = word.bounding_box
        perturbed.append(replace(word, bounding_box=BoundingBox(box.x0 + delta, box.y0, box.x1 + delta, box.y1)))

    grid = reconstruct_table_grid(perturbed, debug=True)

    assert (grid.matrix(), sorted(_span_key(span) for span in grid.merged_ranges())) == reference_signature


def test_span_06_xlsx_contains_exact_mixed_merge_range(tmp_path: Path) -> None:
    result = reconstruct_from_words_file(ROOT / "outputs/extracted_tables/span_06_page_001_words.json", output_dir=tmp_path, debug=True)
    workbook = load_workbook(result.xlsx_path, read_only=False)

    assert "A1:B3" in {str(merged_range) for merged_range in workbook["Table"].merged_cells.ranges}
    assert "Mixed Span Candidates" in workbook.sheetnames
    assert "Mixed Span Decisions" in workbook.sheetnames
    assert "Mixed Span Hierarchy" in workbook.sheetnames
