from __future__ import annotations

import csv
import json
from pathlib import Path

import fitz
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw_pdfs"
GT = ROOT / "data" / "ground_truth"
IMG = ROOT / "data" / "table_images"
MANIFEST = ROOT / "data" / "benchmark_manifest.csv"


def main() -> None:
    RAW.mkdir(parents=True, exist_ok=True)
    GT.mkdir(parents=True, exist_ok=True)
    IMG.mkdir(parents=True, exist_ok=True)

    specs = [
        {
            "table_id": "benchmark_01",
            "category": "simple financial table",
            "notes": "Compact financial table with percentages and parentheses.",
            "bbox": (54, 72, 360, 174),
            "columns": [60, 170, 230, 300],
            "rows": [
                ["Item", "2023", "2024", "Change"],
                ["Revenue", "1000", "1200", "20%"],
                ["Costs", "(600)", "(700)", "16.7%"],
                ["Profit", "400", "500", "25%"],
            ],
            "font": 10,
        },
        {
            "table_id": "benchmark_02",
            "category": "regression table",
            "notes": "Coefficients, parenthetical standard errors, stars, and observations.",
            "bbox": (54, 72, 430, 238),
            "columns": [60, 185, 265, 345],
            "rows": [
                ["Variable", "Model 1", "Model 2", "Model 3"],
                ["Treatment", "0.125**", "0.118**", "0.097*"],
                ["", "(0.041)", "(0.039)", "(0.052)"],
                ["Baseline score", "0.332***", "0.301***", "0.287***"],
                ["", "(0.072)", "(0.068)", "(0.066)"],
                ["Controls", "No", "Yes", "Yes"],
                ["Observations", "1200", "1200", "1200"],
            ],
            "font": 9,
        },
        {
            "table_id": "benchmark_03",
            "category": "correlation matrix",
            "notes": "Lower-triangular matrix with intentional empty upper-triangle cells.",
            "bbox": (54, 72, 410, 198),
            "columns": [60, 145, 205, 265, 325],
            "rows": [
                ["", "Sales", "Costs", "Profit", "Assets"],
                ["Sales", "1.00", "", "", ""],
                ["Costs", "0.82", "1.00", "", ""],
                ["Profit", "0.41", "0.12", "1.00", ""],
                ["Assets", "0.55", "0.49", "0.38", "1.00"],
            ],
            "font": 9,
        },
        {
            "table_id": "benchmark_04",
            "category": "multi-level spanning headers",
            "notes": "Two-level header with adjacent column spans and a row-spanning stub label.",
            "bbox": (54, 72, 430, 198),
            "columns": [60, 165, 225, 300, 360],
            "rows": [
                ["Category", "Q1 Results", "", "Q2 Results", ""],
                ["", "Revenue", "Cost", "Revenue", "Cost"],
                ["Consumer Products", "100", "70", "120", "80"],
                ["", "90", "60", "110", "75"],
                ["", "190", "130", "230", "155"],
            ],
            "merged_ranges": [
                {"row_index": 0, "column_index": 1, "row_span": 1, "column_span": 2},
                {"row_index": 0, "column_index": 3, "row_span": 1, "column_span": 2},
                {"row_index": 2, "column_index": 0, "row_span": 3, "column_span": 1},
            ],
            "header_rows": 2,
            "body_rows": 3,
            "font": 9,
        },
        {
            "table_id": "benchmark_05",
            "category": "borderless table with footnote",
            "notes": "Borderless text table; bbox includes footnote beneath table to measure intrusion.",
            "bbox": (54, 72, 410, 212),
            "columns": [60, 180, 260, 330],
            "rows": [
                ["Measure", "Group A", "Group B", "Diff."],
                ["Completion rate", "91%", "88%", "3%"],
                ["Median days", "14", "17", "-3"],
                ["Satisfaction", "4.2", "4.0", "0.2"],
            ],
            "extra_pdf_lines": [(60, 186, "Note: Values are descriptive and unweighted.")],
            "font": 10,
        },
    ]

    manifest_rows = []
    for index, spec in enumerate(specs, start=1):
        pdf_path = RAW / f"benchmark_{index:02d}.pdf"
        gt_suffix = "xlsx" if spec.get("merged_ranges") else "csv"
        gt_path = GT / f"benchmark_{index:02d}_ground_truth.{gt_suffix}"
        gt_json_path = GT / f"benchmark_{index:02d}_ground_truth.json"
        write_pdf(pdf_path, spec)
        if spec.get("merged_ranges"):
            write_xlsx(gt_path, spec["rows"], spec["merged_ranges"])
            write_structured_json(gt_json_path, spec)
        else:
            write_csv(gt_path, spec["rows"])
        render_crop(pdf_path, IMG / f"benchmark_{index:02d}_crop.png", spec["bbox"])
        manifest_rows.append(
            {
                "table_id": f"benchmark_{index:02d}",
                "pdf_path": f"data/raw_pdfs/benchmark_{index:02d}.pdf",
                "page_number": 1,
                "bbox": ",".join(str(value) for value in spec["bbox"]),
                "category": spec["category"],
                "ground_truth_path": f"data/ground_truth/benchmark_{index:02d}_ground_truth.csv",
                "ground_truth_json_path": "",
                "expected_rows": len(spec["rows"]),
                "expected_columns": max(len(row) for row in spec["rows"]),
                "difficulty_notes": spec["notes"],
                "born_digital": "true",
                "manually_verified": "true",
            }
        )
        if spec.get("merged_ranges"):
            manifest_rows[-1]["ground_truth_path"] = f"data/ground_truth/benchmark_{index:02d}_ground_truth.xlsx"
            manifest_rows[-1]["ground_truth_json_path"] = f"data/ground_truth/benchmark_{index:02d}_ground_truth.json"

    with MANIFEST.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(manifest_rows[0].keys()))
        writer.writeheader()
        writer.writerows(manifest_rows)


def write_pdf(path: Path, spec: dict) -> None:
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((60, 42), spec["category"].title(), fontsize=12)
    y = 82
    row_gap = 18
    for row in spec["rows"]:
        for text, x in zip(row, spec["columns"]):
            if text:
                page.insert_text((x, y), text, fontsize=spec["font"])
        y += row_gap
    for x, y_extra, text in spec.get("extra_pdf_lines", []):
        page.insert_text((x, y_extra), text, fontsize=8)
    doc.save(path)
    doc.close()


def write_csv(path: Path, rows: list[list[str]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows(rows)


def write_xlsx(path: Path, rows: list[list[str]], merged_ranges: list[dict]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "GroundTruth"
    for row_index, row in enumerate(rows, start=1):
        for column_index, text in enumerate(row, start=1):
            worksheet.cell(row_index, column_index, text)
    for merged_range in merged_ranges:
        worksheet.merge_cells(
            start_row=merged_range["row_index"] + 1,
            start_column=merged_range["column_index"] + 1,
            end_row=merged_range["row_index"] + merged_range["row_span"],
            end_column=merged_range["column_index"] + merged_range["column_span"],
        )
    workbook.save(path)


def write_structured_json(path: Path, spec: dict) -> None:
    merged_lookup = {
        (span["row_index"], span["column_index"]): span for span in spec.get("merged_ranges", [])
    }
    covered = {
        (r, c)
        for span in spec.get("merged_ranges", [])
        for r in range(span["row_index"], span["row_index"] + span["row_span"])
        for c in range(span["column_index"], span["column_index"] + span["column_span"])
        if (r, c) != (span["row_index"], span["column_index"])
    }
    logical_cells = []
    for row_index, row in enumerate(spec["rows"]):
        for column_index, text in enumerate(row):
            if (row_index, column_index) in covered:
                continue
            span = merged_lookup.get(
                (row_index, column_index),
                {"row_index": row_index, "column_index": column_index, "row_span": 1, "column_span": 1},
            )
            logical_cells.append(
                {
                    "row_index": row_index,
                    "column_index": column_index,
                    "text": text,
                    "row_span": span["row_span"],
                    "column_span": span["column_span"],
                    "covered_coordinates": [
                        [r, c]
                        for r in range(row_index, row_index + span["row_span"])
                        for c in range(column_index, column_index + span["column_span"])
                    ],
                    "is_header": row_index < spec.get("header_rows", 0),
                    "is_stub": column_index == 0 and row_index >= spec.get("header_rows", 0),
                    "bounding_box": None,
                }
            )
    payload = {
        "physical_rows": len(spec["rows"]),
        "physical_columns": max(len(row) for row in spec["rows"]),
        "header_rows": spec.get("header_rows", 0),
        "body_rows": spec.get("body_rows", len(spec["rows"]) - spec.get("header_rows", 0)),
        "logical_cells": logical_cells,
        "logical_cell_count": len(logical_cells),
        "merged_ranges": [
            cell for cell in logical_cells if cell["row_span"] > 1 or cell["column_span"] > 1
        ],
    }
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def render_crop(pdf_path: Path, output_path: Path, bbox: tuple[int, int, int, int]) -> None:
    doc = fitz.open(pdf_path)
    page = doc.load_page(0)
    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=fitz.Rect(*bbox), alpha=False)
    pix.save(output_path)
    doc.close()


if __name__ == "__main__":
    main()
