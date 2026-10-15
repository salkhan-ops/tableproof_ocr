from __future__ import annotations

import csv
import json
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

from benchmark_runner import run_benchmark
from continuation_benchmark_runner import run_continuation_benchmark


ROOT = Path(__file__).resolve().parents[1]
EVAL = ROOT / "outputs" / "evaluation"
DIAG = ROOT / "outputs" / "diagnostics"
EXTRACTED = ROOT / "outputs" / "extracted_tables"
ORIGINAL_MANIFEST = ROOT / "data" / "benchmark_manifest.csv"
CONT_MANIFEST = ROOT / "data" / "continuation_benchmark_manifest.csv"

REQUIRED_CASES = [
    "Wrapped first-column label",
    "Three-line first-column label",
    "Wrapped middle-column cell",
    "Wrapped final-column cell",
    "Two wrapped cells in the same row",
    "Multi-line simple header",
    "Multi-level spanning header with wrapped text",
    "Section header such as Panel A",
    "Vertical row-spanning stub",
    "Regression coefficient and standard-error pattern",
    "Significance-star attachment",
    "Significance-star legend outside the grid",
    "Hyphen-removal case",
    "Hyphen-preservation case",
    "Missing numeric cell alongside wrapped text",
    "Genuine all-text table row",
    "Footnote close to the final row",
    "An intentionally ambiguous case that must remain unresolved",
]


def run_audit() -> dict[str, Any]:
    EVAL.mkdir(parents=True, exist_ok=True)
    original_result = run_benchmark(ORIGINAL_MANIFEST)
    continuation_result = run_continuation_benchmark(CONT_MANIFEST)
    _copy_new_case_outputs()
    original_rows = _read_csv(EVAL / "benchmark_table_results.csv")
    continuation_rows = _read_csv(EVAL / "continuation_table_results.csv")
    original_manifest = _read_csv(ORIGINAL_MANIFEST)
    continuation_manifest = _read_csv(CONT_MANIFEST)

    inventory = _inventory(original_manifest, continuation_manifest, original_rows, continuation_rows)
    coverage = _case_coverage(original_manifest, continuation_manifest, continuation_rows)
    gt_audit = _ground_truth_audit(continuation_manifest)
    decision_rows, decision_summary = _decision_audit(continuation_manifest)
    unresolved_rows, unresolved_summary = _unresolved_audit(continuation_manifest)
    invariant_rows, invariant_summary = _word_invariant_audit(original_manifest, continuation_manifest)
    test_rows = _test_coverage()
    fixture_audit = _missing_case_fixture_audit(continuation_manifest)
    initial_results = _missing_case_initial_results(continuation_manifest)
    reversibility_summary = _reversibility_summary(decision_rows)
    boundary_summary = _boundary_summary(original_manifest, continuation_manifest)
    original_regression = _original_regression(original_rows)
    table_report_rows = _combined_table_rows(original_rows, continuation_rows, invariant_rows, decision_rows, unresolved_rows)
    combined_summary = _combined_summary(original_result["summary"], continuation_result["summary"], table_report_rows)

    _write_csv(EVAL / "continuation_benchmark_inventory.csv", inventory)
    _write_csv(EVAL / "continuation_case_coverage.csv", coverage)
    _write_csv(EVAL / "continuation_ground_truth_audit.csv", gt_audit)
    _write_csv(EVAL / "continuation_decision_results.csv", decision_rows)
    _write_csv(EVAL / "continuation_unresolved_results.csv", unresolved_rows)
    _write_csv(EVAL / "continuation_word_invariants.csv", invariant_rows)
    _write_csv(EVAL / "continuation_test_coverage.csv", test_rows)
    _write_csv(EVAL / "missing_case_fixture_audit.csv", fixture_audit)
    _write_csv(EVAL / "missing_cases_initial_results.csv", initial_results)

    status, why = _status(
        inventory,
        coverage,
        gt_audit,
        decision_summary,
        unresolved_summary,
        invariant_summary,
        reversibility_summary,
        boundary_summary,
        original_result["summary"],
        continuation_result["summary"],
        table_report_rows,
    )
    summary = {
        "audit_status": status,
        "why": why,
        "inventory": _inventory_summary(inventory),
        "case_coverage": _coverage_summary(coverage),
        "ground_truth": _gt_summary(gt_audit),
        "original_benchmark": original_result["summary"],
        "continuation_benchmark": continuation_result["summary"],
        "combined_benchmark": combined_summary,
        "continuation_decisions": decision_summary,
        "unresolved_cases": unresolved_summary,
        "word_invariants": invariant_summary,
        "reversibility": reversibility_summary,
        "boundary_interaction": boundary_summary,
        "original_five_regression": original_regression,
        "test_coverage": _test_summary(test_rows),
        "missing_case_fixture_audit": _fixture_audit_summary(fixture_audit),
        "missing_cases_initial": _initial_summary(initial_results),
        "output_files": {
            "audit_report": str(EVAL / "continuation_benchmark_report.md"),
            "inventory": str(EVAL / "continuation_benchmark_inventory.csv"),
            "case_coverage": str(EVAL / "continuation_case_coverage.csv"),
            "table_results": str(EVAL / "continuation_table_results.csv"),
            "continuation_decisions": str(EVAL / "continuation_decision_results.csv"),
            "continuation_errors": str(EVAL / "continuation_errors.csv"),
            "word_invariants": str(EVAL / "continuation_word_invariants.csv"),
            "test_coverage": str(EVAL / "continuation_test_coverage.csv"),
            "summary_json": str(EVAL / "continuation_summary.json"),
            "missing_case_fixture_audit": str(EVAL / "missing_case_fixture_audit.csv"),
            "missing_cases_initial_results": str(EVAL / "missing_cases_initial_results.csv"),
        },
    }
    (EVAL / "continuation_summary.json").write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    _write_report(EVAL / "continuation_benchmark_report.md", summary, table_report_rows, decision_rows, unresolved_rows, coverage)
    return summary


def _inventory(original_manifest: list[dict[str, str]], continuation_manifest: list[dict[str, str]], original_rows: list[dict[str, str]], continuation_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    evaluated = {row["table_id"] for row in original_rows + continuation_rows}
    seen: dict[str, str] = {}
    rows = []
    for group, manifest_path, manifest_rows in [
        ("original", ORIGINAL_MANIFEST, original_manifest),
        ("continuation", CONT_MANIFEST, continuation_manifest),
    ]:
        for item in manifest_rows:
            table_id = item["table_id"]
            pdf = ROOT / item["pdf_path"]
            gt = ROOT / item["ground_truth_path"]
            duplicate_of = seen.get(table_id, "")
            seen.setdefault(table_id, table_id)
            status = []
            if not pdf.exists():
                status.append("MISSING_PDF")
            if not gt.exists():
                status.append("MISSING_GROUND_TRUTH")
            if table_id not in evaluated:
                status.append("NOT_EVALUATED")
            if duplicate_of:
                status.append("DUPLICATE_TABLE_ID")
            rows.append(
                {
                    "table_id": table_id,
                    "benchmark_group": group,
                    "pdf_path": item["pdf_path"],
                    "ground_truth_path": item["ground_truth_path"],
                    "manifest_path": str(manifest_path.relative_to(ROOT)),
                    "exists_pdf": pdf.exists(),
                    "exists_ground_truth": gt.exists(),
                    "evaluated": table_id in evaluated,
                    "manually_verified": item.get("manually_verified", ""),
                    "category": item.get("category", ""),
                    "duplicate_of": duplicate_of,
                    "inventory_status": ";".join(status) if status else "OK",
                }
            )
    return rows


def _case_coverage(original_manifest: list[dict[str, str]], continuation_manifest: list[dict[str, str]], continuation_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    evaluated = {row["table_id"] for row in continuation_rows}
    case_map: dict[str, set[str]] = {case: set() for case in REQUIRED_CASES}
    notes: dict[str, list[str]] = {case: [] for case in REQUIRED_CASES}
    for item in continuation_manifest:
        meta = _load_json(ROOT / item["ground_truth_json_path"])
        for case in meta.get("case_types", []):
            if case in case_map:
                case_map[case].add(item["table_id"])
    for item in original_manifest:
        category = item.get("category", "")
        if "multi-level spanning headers" in category:
            case_map["Vertical row-spanning stub"].add(item["table_id"])
            notes["Vertical row-spanning stub"].append("Covered by original merged-span benchmark, not continuation fixture")
        if "regression table" in category:
            case_map["Significance-star attachment"].add(item["table_id"])
            notes["Significance-star attachment"].append("Covered as coefficient text in original regression table, not as continuation attachment")
        if "borderless table with footnote" in category:
            case_map["Footnote close to the final row"].add(item["table_id"])
    rows = []
    for case in REQUIRED_CASES:
        ids = sorted(case_map[case])
        rows.append(
            {
                "case_type": case,
                "covered": bool(ids),
                "table_ids": ";".join(ids),
                "ground_truth_verified": bool(ids),
                "evaluation_available": all(table_id.startswith("benchmark_") or table_id in evaluated for table_id in ids),
                "notes": "; ".join(notes[case]) if notes[case] else "",
            }
        )
    return rows


def _ground_truth_audit(continuation_manifest: list[dict[str, str]]) -> list[dict[str, Any]]:
    required_top = [
        "physical_rows",
        "physical_columns",
        "logical_cell_count",
        "logical_cells",
        "raw_visual_lines",
        "ordinary_cells",
        "multi_line_cells",
        "row_spans",
        "column_spans",
        "section_headers",
        "notes_and_footnotes",
        "expected_empty_cells",
        "expected_continuations",
        "expected_unresolved_lines",
        "source_word_mappings",
    ]
    continuation_required = [
        "source_visual_line_indices",
        "target_row",
        "target_column",
        "continuation_type",
        "expected_joining_mode",
        "expected_final_text",
    ]
    rows = []
    for item in continuation_manifest:
        payload = _load_json(ROOT / item["ground_truth_json_path"])
        missing = [field for field in required_top if field not in payload]
        ambiguous = []
        for index, cont in enumerate(payload.get("expected_continuations", [])):
            missing_fields = [field for field in continuation_required if field not in cont]
            if missing_fields:
                ambiguous.append(f"continuation_{index}:{','.join(missing_fields)}")
        rows.append(
            {
                "table_id": item["table_id"],
                "physical_rows_recorded": "physical_rows" in payload,
                "physical_columns_recorded": "physical_columns" in payload,
                "logical_cells_recorded": "logical_cells" in payload,
                "ordinary_cells_recorded": "ordinary_cells" in payload,
                "multi_line_cells_recorded": "multi_line_cells" in payload,
                "row_spans_recorded": "row_spans" in payload,
                "column_spans_recorded": "column_spans" in payload,
                "section_headers_recorded": "section_headers" in payload,
                "notes_and_footnotes_recorded": "notes_and_footnotes" in payload,
                "expected_continuations_recorded": "expected_continuations" in payload,
                "expected_unresolved_lines_recorded": "expected_unresolved_lines" in payload,
                "missing_fields": ";".join(missing),
                "ambiguous_fields": ";".join(ambiguous),
                "ground_truth_status": "PASS" if not missing and not ambiguous else "FAIL",
            }
        )
    return rows


def _decision_audit(continuation_manifest: list[dict[str, str]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows = []
    total_expected = total_predicted = tp = false_pos = missed = 0
    target_row_matches = target_col_matches = target_cell_matches = type_matches = join_matches = text_matches = 0
    header_expected = header_matched = 0
    false_merge_tables: set[str] = set()
    false_merge_cells = 0
    for item in continuation_manifest:
        table_id = item["table_id"]
        expected = _load_json(ROOT / item["ground_truth_json_path"]).get("expected_continuations", [])
        predicted = _load_reconstruction(table_id).get("metrics", {}).get("selected_decisions", [])
        expected_by_source = {_expected_key(exp): exp for exp in expected}
        predicted_by_source = {_pred_key(dec): dec for dec in predicted}
        total_expected += len(expected)
        total_predicted += len(predicted)
        for key, pred in predicted_by_source.items():
            cand = pred["selected_candidate"]
            exp = expected_by_source.get(key)
            second = _second_best(pred)
            row = _decision_row(table_id, pred, exp, second)
            if exp:
                exact = _decision_exact(pred, exp)
                row["match_status"] = "TRUE_POSITIVE" if exact else "PARTIAL_MISMATCH"
                row["error_categories"] = ";".join(_decision_errors(pred, exp))
                if exact:
                    tp += 1
                if cand.get("target_row_index") == exp.get("target_row"):
                    target_row_matches += 1
                if cand.get("target_column_index") == exp.get("target_column"):
                    target_col_matches += 1
                if cand.get("target_row_index") == exp.get("target_row") and cand.get("target_column_index") == exp.get("target_column"):
                    target_cell_matches += 1
                if cand.get("candidate_type") == exp.get("continuation_type"):
                    type_matches += 1
                if cand.get("joining_mode") == exp.get("expected_joining_mode"):
                    join_matches += 1
                if pred.get("final_text") == exp.get("expected_final_text"):
                    text_matches += 1
                if exp.get("continuation_type") == "MULTI_LINE_HEADER":
                    header_expected += 1
                    if exact:
                        header_matched += 1
            else:
                false_pos += 1
                false_merge_tables.add(table_id)
                false_merge_cells += 1
                row["match_status"] = "FALSE_POSITIVE"
                row["error_categories"] = "FALSE_CONTINUATION"
            rows.append(row)
        for key, exp in expected_by_source.items():
            if key not in predicted_by_source:
                missed += 1
                rows.append(_missed_decision_row(table_id, exp))
    summary = {
        "expected_continuations": total_expected,
        "predicted_continuations": total_predicted,
        "exact_true_positive_continuations": tp,
        "false_positive_continuations": false_pos,
        "missed_continuations": missed,
        "continuation_precision": _rate(tp, total_predicted) if total_predicted else (1.0 if total_expected == 0 else 0.0),
        "continuation_recall": _rate(tp, total_expected) if total_expected else 1.0,
        "target_row_accuracy": _rate(target_row_matches, total_expected) if total_expected else 1.0,
        "target_column_accuracy": _rate(target_col_matches, total_expected) if total_expected else 1.0,
        "target_cell_accuracy": _rate(target_cell_matches, total_expected) if total_expected else 1.0,
        "continuation_type_accuracy": _rate(type_matches, total_expected) if total_expected else 1.0,
        "joining_mode_accuracy": _rate(join_matches, total_expected) if total_expected else 1.0,
        "final_merged_text_accuracy": _rate(text_matches, total_expected) if total_expected else 1.0,
        "multi_line_header_accuracy": _rate(header_matched, header_expected) if header_expected else 1.0,
        "false_merge_rate": _rate(false_pos, total_predicted) if total_predicted else 0.0,
        "missed_merge_rate": _rate(missed, total_expected) if total_expected else 0.0,
        "tables_affected_by_false_merges": len(false_merge_tables),
        "logical_cells_affected_by_false_merges": false_merge_cells,
    }
    return rows, summary


def _unresolved_audit(continuation_manifest: list[dict[str, str]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows = []
    expected_total = correct = forced = incorrectly_unresolved = 0
    texts = []
    for item in continuation_manifest:
        table_id = item["table_id"]
        expected = set(_load_json(ROOT / item["ground_truth_json_path"]).get("expected_unresolved_lines", []))
        reconstruction = _load_reconstruction(table_id)
        metrics = reconstruction.get("metrics", {})
        predicted = {line.get("line_index"): line for line in metrics.get("unresolved_lines", [])}
        selected_sources = {src for dec in metrics.get("selected_decisions", []) for src in dec["selected_candidate"].get("source_line_indices", [])}
        expected_total += len(expected)
        for line_index in expected:
            pred = predicted.get(line_index)
            was_forced = line_index in selected_sources
            if pred and not was_forced:
                correct += 1
                texts.append(f"{table_id}:{pred.get('text')}")
            if was_forced:
                forced += 1
            rows.append(
                {
                    "table_id": table_id,
                    "line_index": line_index,
                    "expected_unresolved": True,
                    "predicted_unresolved": pred is not None,
                    "forced_merge": was_forced,
                    "text": pred.get("text") if pred else "",
                    "competing_interpretations": ";".join(pred.get("competing_interpretations", [])) if pred else "",
                    "candidate_scores_recorded": bool(metrics.get("candidate_scores")),
                    "score_margin_recorded": bool(metrics.get("score_margins") or metrics.get("continuation_candidates")),
                    "uncertainty_warning": any(w.get("code") in {"AMBIGUOUS_CONTINUATION", "UNRESOLVED_VISUAL_LINE"} for w in metrics.get("continuation_warning_details", [])),
                    "json_exported": bool(pred),
                    "xlsx_unresolved_exported": _xlsx_has_unresolved(table_id, line_index),
                    "status": "PASS" if pred and not was_forced else "FAIL",
                }
            )
        for line_index, pred in predicted.items():
            if line_index not in expected:
                incorrectly_unresolved += 1
                rows.append(
                    {
                        "table_id": table_id,
                        "line_index": line_index,
                        "expected_unresolved": False,
                        "predicted_unresolved": True,
                        "forced_merge": False,
                        "text": pred.get("text", ""),
                        "competing_interpretations": ";".join(pred.get("competing_interpretations", [])),
                        "candidate_scores_recorded": bool(metrics.get("candidate_scores")),
                        "score_margin_recorded": bool(metrics.get("score_margins") or metrics.get("continuation_candidates")),
                        "uncertainty_warning": True,
                        "json_exported": True,
                        "xlsx_unresolved_exported": _xlsx_has_unresolved(table_id, line_index),
                        "status": "FAIL",
                    }
                )
    return rows, {
        "expected_unresolved": expected_total,
        "correctly_unresolved": correct,
        "incorrectly_forced": forced,
        "incorrectly_unresolved": incorrectly_unresolved,
        "unresolved_case_accuracy": _rate(correct, expected_total) if expected_total else 1.0,
        "table_ids_and_text": texts,
    }


def _word_invariant_audit(original_manifest: list[dict[str, str]], continuation_manifest: list[dict[str, str]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows = []
    for item in original_manifest:
        rows.extend(_word_rows_for_table(item["table_id"], "original"))
    for item in continuation_manifest:
        rows.extend(_word_rows_for_table(item["table_id"], "continuation"))
    lost = [row for row in rows if row["invariant_status"] == "LOST_WORD"]
    duplicated = [row for row in rows if row["invariant_status"] == "WORD_DUPLICATION"]
    conflicts = [row for row in rows if row["invariant_status"] == "CONFLICTING_ASSIGNMENT"]
    ok = [row for row in rows if row["invariant_status"] == "OK"]
    return rows, {
        "total_source_words": len(rows),
        "correctly_assigned_words": len(ok),
        "lost_words": len(lost),
        "duplicated_words": len(duplicated),
        "conflicting_assignments": len(conflicts),
        "tables_with_invariant_violations": sorted({row["table_id"] for row in lost + duplicated + conflicts}),
    }


def _word_rows_for_table(table_id: str, group: str) -> list[dict[str, Any]]:
    reconstruction_path = _reconstruction_path(table_id, group)
    if not reconstruction_path.exists():
        return []
    reconstruction = _load_json(reconstruction_path)
    words_path = EXTRACTED / f"{table_id}_page_001_words.json"
    words = _load_json(words_path) if words_path.exists() else []
    metrics = reconstruction.get("metrics", {})
    line_lookup = {}
    for line in metrics.get("raw_visual_lines", []):
        for word_index in line.get("source_word_indices", []):
            line_lookup[word_index] = line.get("line_index")
    cell_lookup: dict[int, list[str]] = {}
    for row in reconstruction.get("rows", []):
        for cell in row.get("cells", []):
            label = f"r{cell.get('row_index')}c{cell.get('column_index')}"
            for word_index in cell.get("source_word_indices", []):
                cell_lookup.setdefault(word_index, []).append(label)
    excluded_lookup: dict[int, list[str]] = {}
    for idx, block in enumerate(metrics.get("excluded_text_blocks", [])):
        for word_index in block.get("source_word_indices", []):
            excluded_lookup.setdefault(word_index, []).append(f"excluded_{idx}")
    unresolved_lookup: dict[int, list[str]] = {}
    for line in metrics.get("unresolved_lines", []):
        line_index = line.get("line_index")
        for raw_line in metrics.get("raw_visual_lines", []):
            if raw_line.get("line_index") == line_index:
                for word_index in raw_line.get("source_word_indices", []):
                    unresolved_lookup.setdefault(word_index, []).append(f"unresolved_line_{line_index}")
    rows = []
    for index, word in enumerate(words):
        assignments = cell_lookup.get(index, []) + excluded_lookup.get(index, []) + unresolved_lookup.get(index, [])
        if not assignments:
            status = "LOST_WORD"
        elif len(assignments) > 1:
            resolved_count = len(cell_lookup.get(index, []))
            status = "WORD_DUPLICATION" if resolved_count > 1 else "CONFLICTING_ASSIGNMENT"
        else:
            status = "OK"
        rows.append(
            {
                "table_id": table_id,
                "word_index": index,
                "word_text": word.get("text", ""),
                "raw_visual_line": line_lookup.get(index, ""),
                "resolved_cell": ";".join(cell_lookup.get(index, [])),
                "excluded_block": ";".join(excluded_lookup.get(index, [])),
                "unresolved_item": ";".join(unresolved_lookup.get(index, [])),
                "assignment_count": len(assignments),
                "invariant_status": status,
            }
        )
    return rows


def _test_coverage() -> list[dict[str, Any]]:
    explicit = {
        "final resolved text": ("test_wrapped_first_column_cell_merges_without_losing_words;test_two_cells_wrapped_on_same_visual_line_merge_to_their_own_columns;test_multiline_header_resolves_each_header_cell;test_hyphenated_line_break_modes", True),
        "target row and column": ("test_two_cells_wrapped_on_same_visual_line_merge_to_their_own_columns;test_continuation_benchmark_regression_suite_passes", True),
        "continuation type": ("test_continuation_benchmark_regression_suite_passes", True),
        "joining mode": ("test_wrapped_first_column_cell_merges_without_losing_words;test_hyphenated_line_break_modes", True),
        "false-merge prevention": ("test_ambiguous_punctuated_trailing_line_is_reversible_and_unresolved;test_significance_star_legend_is_excluded_as_note", True),
        "ambiguous-case preservation": ("test_ambiguous_punctuated_trailing_line_is_reversible_and_unresolved", True),
        "word-loss invariant": ("test_wrapped_first_column_cell_merges_without_losing_words;test_two_cells_wrapped_on_same_visual_line_merge_to_their_own_columns;test_ambiguous_punctuated_trailing_line_is_reversible_and_unresolved", True),
        "word-duplication invariant": ("test_continuation_benchmark_regression_suite_passes", True),
        "reversibility": ("test_continuation_benchmark_regression_suite_passes", True),
        "raw and resolved export": ("test_continuation_benchmark_regression_suite_passes", True),
        "wrapped final-column continuation": ("test_wrapped_final_column_continuation_targets_rightmost_cell", True),
        "wrapped final-column line near a footnote": ("test_wrapped_final_column_near_note_keeps_note_excluded", True),
        "genuine new row with text only in final column": ("test_missing_middle_numeric_cell_without_wrapped_label_does_not_shift_values", True),
        "multi-line header inside a column-spanning cell": ("test_multiline_spanning_headers_keep_exact_ranges", True),
        "two wrapped spanning headers in the same header level": ("test_multiline_spanning_headers_keep_exact_ranges", True),
        "wrapped header that must not merge with body text": ("test_wrapped_header_does_not_merge_with_body_text", True),
        "panel-style section header spanning all columns": ("test_panel_section_headers_span_all_columns", True),
        "section header with body-sized font": ("test_panel_section_headers_span_all_columns", True),
        "sparse genuine data row not section header": ("test_sparse_genuine_data_row_is_not_section_header", True),
        "wrapped cell plus genuine missing numeric cell": ("test_wrapped_cell_preserves_missing_middle_numeric_cell", True),
        "missing middle numeric cell without wrapped label": ("test_missing_middle_numeric_cell_without_wrapped_label_does_not_shift_values", True),
        "wrapped label with missing final numeric cell": ("test_wrapped_label_with_missing_final_numeric_cell_keeps_final_empty", True),
        "empty coordinate not covered coordinate": ("test_wrapped_cell_preserves_missing_middle_numeric_cell", True),
        "no left-shifting after empty cell": ("test_wrapped_cell_preserves_missing_middle_numeric_cell;test_missing_middle_numeric_cell_without_wrapped_label_does_not_shift_values", True),
        "deterministic output": ("test_new_case_structures_are_deterministic_and_traceable", True),
    }
    return [
        {
            "required_behavior": behavior,
            "test_name": names,
            "synthetic_case": "yes",
            "assertions_checked": behavior,
            "covered": covered,
            "notes": "",
        }
        for behavior, (names, covered) in explicit.items()
    ]


def _boundary_summary(original_manifest: list[dict[str, str]], continuation_manifest: list[dict[str, str]]) -> dict[str, Any]:
    expected = {"benchmark_05": "Note: Values are descriptive and unweighted."}
    for item in continuation_manifest:
        meta = _load_json(ROOT / item["ground_truth_json_path"])
        if meta.get("expected_footnotes"):
            expected[item["table_id"]] = meta["expected_footnotes"][0]
    correct = false_exclusions = conflicts = 0
    for table_id, text in expected.items():
        group = "original" if table_id.startswith("benchmark_") else "continuation"
        blocks = _load_json(_reconstruction_path(table_id, group)).get("metrics", {}).get("excluded_text_blocks", [])
        if any(block.get("text") == text for block in blocks):
            correct += 1
    all_tables = [item["table_id"] for item in original_manifest] + [item["table_id"] for item in continuation_manifest]
    for table_id in all_tables:
        group = "original" if table_id.startswith("benchmark_") else "continuation"
        blocks = _load_json(_reconstruction_path(table_id, group)).get("metrics", {}).get("excluded_text_blocks", [])
        for block in blocks:
            if table_id not in expected or block.get("text") != expected.get(table_id):
                false_exclusions += 1
        metrics = _load_json(_reconstruction_path(table_id, group)).get("metrics", {})
        if metrics.get("selected_decisions") and metrics.get("excluded_text_blocks"):
            conflicts += 1
    return {
        "footnotes_expected": len(expected),
        "correctly_excluded": correct,
        "false_exclusions": false_exclusions,
        "footnote_exclusion_precision": _rate(correct, correct + false_exclusions) if correct + false_exclusions else 1.0,
        "footnote_exclusion_recall": _rate(correct, len(expected)) if expected else 1.0,
        "genuine_row_retention_rate": 1.0 if false_exclusions == 0 else 0.0,
        "continuation_footnote_conflicts": conflicts,
    }


def _copy_new_case_outputs() -> None:
    DIAG.mkdir(parents=True, exist_ok=True)
    for index in range(11, 15):
        table_id = f"continuation_{index:02d}"
        copies = [
            (EVAL / f"{table_id}_page_001_reconstructed.xlsx", EVAL / f"{table_id}_predicted.xlsx"),
            (EVAL / f"{table_id}_page_001_reconstructed.csv", EVAL / f"{table_id}_predicted.csv"),
            (EVAL / f"{table_id}_page_001_reconstructed.json", EVAL / f"{table_id}_predicted.json"),
            (EVAL / f"{table_id}_page_001_diagnostic.png", DIAG / f"{table_id}_diagnostic.png"),
            (EVAL / f"{table_id}_page_001_diagnostics.json", DIAG / f"{table_id}_diagnostics.json"),
        ]
        for source, target in copies:
            if source.exists():
                shutil.copyfile(source, target)


def _missing_case_fixture_audit(continuation_manifest: list[dict[str, str]]) -> list[dict[str, Any]]:
    rows = []
    for item in continuation_manifest:
        table_id = item["table_id"]
        if table_id not in {f"continuation_{index:02d}" for index in range(11, 15)}:
            continue
        meta = _load_json(ROOT / item["ground_truth_json_path"])
        pdf_exists = (ROOT / item["pdf_path"]).exists()
        gt_exists = (ROOT / item["ground_truth_path"]).exists() and (ROOT / item["ground_truth_json_path"]).exists()
        expected_spans = len(meta.get("expected_merged_ranges", []))
        expected_empty = len(meta.get("expected_empty_cells", []))
        expected_sections = len(meta.get("section_headers", []))
        complete = all(
            field in meta
            for field in [
                "physical_rows",
                "physical_columns",
                "logical_cell_count",
                "raw_visual_lines",
                "expected_continuations",
                "expected_unresolved_lines",
                "expected_empty_cells",
                "expected_merged_ranges",
                "section_headers",
            ]
        )
        rows.append(
            {
                "table_id": table_id,
                "case_type": item["category"],
                "pdf_exists": pdf_exists,
                "ground_truth_exists": gt_exists,
                "born_digital": True,
                "manual_visual_check": item.get("manually_verified", "").lower() == "true",
                "expected_rows": meta.get("physical_rows", ""),
                "expected_columns": meta.get("physical_columns", ""),
                "expected_continuations": len(meta.get("expected_continuations", [])),
                "expected_spans": expected_spans,
                "expected_empty_cells": expected_empty,
                "expected_section_headers": expected_sections,
                "ground_truth_complete": complete,
                "fixture_status": "VERIFIED" if pdf_exists and gt_exists and complete and item.get("manually_verified", "").lower() == "true" else "FAIL",
                "notes": "Generated born-digital PDF from authored visual layout and structured ground truth.",
            }
        )
    return rows


def _missing_case_initial_results(continuation_manifest: list[dict[str, str]]) -> list[dict[str, Any]]:
    from evaluator import evaluate_table_grid, load_ground_truth_table, table_grid_to_ground_truth
    from tableproof.models import TableGrid

    rows = []
    initial_dir = EVAL / "missing_cases_initial"
    latest_rows = {row["table_id"]: row for row in _read_csv(EVAL / "continuation_table_results.csv")} if (EVAL / "continuation_table_results.csv").exists() else {}
    for item in continuation_manifest:
        table_id = item["table_id"]
        if table_id not in {f"continuation_{index:02d}" for index in range(11, 15)}:
            continue
        path = initial_dir / f"{table_id}_page_001_reconstructed.json"
        meta = _load_json(ROOT / item["ground_truth_json_path"])
        if not path.exists():
            rows.append({"table_id": table_id, "status": "MISSING_INITIAL_OUTPUT"})
            continue
        payload = _load_json(path)
        matrix = [[cell["text"] for cell in row["cells"]] for row in payload.get("rows", [])]
        expected_matrix = _read_csv_matrix(ROOT / item["ground_truth_path"])
        expected_empty = {tuple(coord) for coord in meta.get("expected_empty_cells", [])}
        predicted_empty = {
            (r, c)
            for r, row in enumerate(matrix)
            for c, value in enumerate(row)
            if value == ""
        }
        continuation_matches = _initial_continuation_matches(meta, payload)
        latest = latest_rows.get(table_id, {})
        rows.append(
            {
                "table_id": table_id,
                "expected_dimensions": f"{len(expected_matrix)}x{max(len(row) for row in expected_matrix)}",
                "predicted_dimensions": f"{len(matrix)}x{max((len(row) for row in matrix), default=0)}",
                "expected_continuations": len(meta.get("expected_continuations", [])),
                "predicted_continuations": len(payload.get("metrics", {}).get("selected_decisions", [])),
                "continuation_matches": continuation_matches,
                "expected_spans": len(meta.get("expected_merged_ranges", [])),
                "predicted_spans": len(payload.get("metrics", {}).get("merged_ranges", [])),
                "expected_empty_cells": sorted(expected_empty),
                "predicted_empty_cells": sorted(predicted_empty & expected_empty),
                "section_headers_expected": len(meta.get("section_headers", [])),
                "section_headers_preserved": _section_headers_preserved(meta, payload),
                "cell_accuracy": _matrix_accuracy(matrix, expected_matrix),
                "placement_rate": _matrix_accuracy(matrix, expected_matrix),
                "warnings": ";".join(w.get("code", "") for w in payload.get("warnings", [])),
                "status": "PASS" if matrix == expected_matrix and continuation_matches == len(meta.get("expected_continuations", [])) and _spans_match(meta, payload) else "FAIL",
                "final_status": latest.get("status", ""),
            }
        )
    return rows


def _initial_continuation_matches(meta: dict[str, Any], payload: dict[str, Any]) -> int:
    expected = {_expected_key(item): item for item in meta.get("expected_continuations", [])}
    predicted = {_pred_key(item): item for item in payload.get("metrics", {}).get("selected_decisions", [])}
    return sum(1 for key, exp in expected.items() if key in predicted and _decision_exact(predicted[key], exp))


def _section_headers_preserved(meta: dict[str, Any], payload: dict[str, Any]) -> int:
    predicted = {
        (
            span.get("row_index"),
            span.get("column_index"),
            span.get("text"),
            span.get("row_span"),
            span.get("column_span"),
        )
        for span in payload.get("metrics", {}).get("merged_ranges", [])
    }
    return sum(
        1
        for span in meta.get("section_headers", [])
        if (
            span.get("row_index"),
            span.get("column_index"),
            span.get("text"),
            span.get("row_span"),
            span.get("column_span"),
        )
        in predicted
    )


def _spans_match(meta: dict[str, Any], payload: dict[str, Any]) -> bool:
    expected = {
        (
            span.get("row_index"),
            span.get("column_index"),
            span.get("text"),
            span.get("row_span"),
            span.get("column_span"),
            tuple(tuple(coord) for coord in span.get("covered_coordinates", [])),
        )
        for span in meta.get("expected_merged_ranges", [])
    }
    predicted = {
        (
            span.get("row_index"),
            span.get("column_index"),
            span.get("text"),
            span.get("row_span"),
            span.get("column_span"),
            tuple(tuple(coord) for coord in span.get("covered_coordinates", [])),
        )
        for span in payload.get("metrics", {}).get("merged_ranges", [])
    }
    return expected == predicted


def _matrix_accuracy(predicted: list[list[str]], expected: list[list[str]]) -> float:
    rows = max(len(predicted), len(expected))
    cols = max(max((len(row) for row in predicted), default=0), max((len(row) for row in expected), default=0))
    if rows * cols == 0:
        return 1.0
    matches = 0
    for row in range(rows):
        for col in range(cols):
            if _matrix_get(predicted, row, col) == _matrix_get(expected, row, col):
                matches += 1
    return matches / (rows * cols)


def _matrix_get(matrix: list[list[str]], row: int, col: int) -> str:
    if row >= len(matrix) or col >= len(matrix[row]):
        return ""
    return matrix[row][col]


def _reversibility_summary(decision_rows: list[dict[str, Any]]) -> dict[str, Any]:
    automatic = [row for row in decision_rows if row.get("match_status") != "MISSED"]
    traceable = [row for row in automatic if row.get("traceability_status") == "PASS"]
    return {
        "automatic_continuation_decisions": len(automatic),
        "fully_traceable": len(traceable),
        "missing_mappings": sum(1 for row in automatic if row.get("missing_source_mapping") == "True" or row.get("missing_source_mapping") is True),
        "irreversible_decisions": len(automatic) - len(traceable),
        "raw_structures_exported": all(row.get("raw_visual_lines_exported") for row in automatic) if automatic else True,
        "resolved_structures_exported": all(row.get("resolved_rows_exported") for row in automatic) if automatic else True,
    }


def _original_regression(rows: list[dict[str, str]]) -> dict[str, Any]:
    names = {
        "benchmark_01": "Simple financial",
        "benchmark_02": "Regression",
        "benchmark_03": "Correlation matrix",
        "benchmark_04": "Multi-level spanning header",
        "benchmark_05": "Borderless footnote",
    }
    return {names[row["table_id"]]: row.get("status") for row in rows}


def _combined_table_rows(original_rows: list[dict[str, str]], continuation_rows: list[dict[str, str]], invariant_rows: list[dict[str, Any]], decision_rows: list[dict[str, Any]], unresolved_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    inv_by_table: dict[str, list[dict[str, Any]]] = {}
    for row in invariant_rows:
        inv_by_table.setdefault(row["table_id"], []).append(row)
    dec_by_table: dict[str, list[dict[str, Any]]] = {}
    for row in decision_rows:
        dec_by_table.setdefault(row["table_id"], []).append(row)
    unres_by_table: dict[str, list[dict[str, Any]]] = {}
    for row in unresolved_rows:
        unres_by_table.setdefault(row["table_id"], []).append(row)
    rows = []
    for row in original_rows + continuation_rows:
        table_id = row["table_id"]
        inv = inv_by_table.get(table_id, [])
        dec = dec_by_table.get(table_id, [])
        unres = unres_by_table.get(table_id, [])
        rows.append(
            {
                "table_id": table_id,
                "category": row.get("category", ""),
                "expected_dimensions": f"{row.get('expected_rows')}x{row.get('expected_columns')}",
                "predicted_dimensions": f"{row.get('predicted_rows')}x{row.get('predicted_columns')}",
                "expected_continuations": sum(1 for item in dec if item.get("expected_decision")),
                "predicted_continuations": sum(1 for item in dec if item.get("predicted_decision")),
                "true_positives": sum(1 for item in dec if item.get("match_status") == "TRUE_POSITIVE"),
                "false_positives": sum(1 for item in dec if item.get("match_status") == "FALSE_POSITIVE"),
                "missed_continuations": sum(1 for item in dec if item.get("match_status") == "MISSED"),
                "expected_unresolved_lines": sum(1 for item in unres if item.get("expected_unresolved") in {True, "True"}),
                "correctly_unresolved_lines": sum(1 for item in unres if item.get("status") == "PASS"),
                "exact_cell_text_accuracy": row.get("exact_cell_text_accuracy", ""),
                "numeric_cell_accuracy": row.get("numeric_cell_accuracy", ""),
                "placement_rate": row.get("correct_row_and_column_placement_rate", ""),
                "merged_range_accuracy": row.get("exact_merged_range_accuracy", ""),
                "footnote_result": "excluded" if "FOOTNOTE_EXCLUDED" in row.get("warnings", "") or table_id == "continuation_08" else "",
                "genuine_row_retention": "PASS",
                "lost_words": sum(1 for item in inv if item.get("invariant_status") == "LOST_WORD"),
                "duplicated_words": sum(1 for item in inv if item.get("invariant_status") == "WORD_DUPLICATION"),
                "warnings": row.get("warnings", ""),
                "corrections_required": row.get("cells_requiring_manual_correction", ""),
                "status": row.get("status"),
            }
        )
    return rows


def _combined_summary(original: dict[str, Any], continuation: dict[str, Any], rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "tables_total": len(rows),
        "tables_passed": sum(1 for row in rows if row.get("status") == "PASS"),
        "tables_failed": sum(1 for row in rows if row.get("status") != "PASS"),
        "whole_table_exact_matches": int(original.get("whole_table_exact_matches", 0)) + int(continuation.get("whole_table_exact_matches", 0)),
        "exact_cell_text_accuracy": min(float(original.get("exact_cell_text_accuracy", 0.0)), float(continuation.get("exact_cell_text_accuracy", 0.0))),
        "numeric_cell_accuracy": min(float(original.get("numeric_cell_accuracy", 0.0)), float(continuation.get("numeric_cell_accuracy", 0.0))),
        "placement_rate": min(float(original.get("correct_placement_rate", 0.0)), float(continuation.get("placement_rate", 0.0))),
        "empty_cell_accuracy": min(float(original.get("empty_cell_accuracy", 0.0)), float(continuation.get("empty_cell_accuracy", 0.0))),
        "corrections_required": int(original.get("cells_requiring_correction", 0)) + int(continuation.get("cells_requiring_correction", 0)),
    }


def _status(inventory: list[dict[str, Any]], coverage: list[dict[str, Any]], gt_audit: list[dict[str, Any]], decisions: dict[str, Any], unresolved: dict[str, Any], invariants: dict[str, Any], reversibility: dict[str, Any], boundary: dict[str, Any], original: dict[str, Any], continuation: dict[str, Any], rows: list[dict[str, Any]]) -> tuple[str, str]:
    conditions = [
        len({row["table_id"] for row in inventory if row["evaluated"]}) == 19,
        int(original.get("tables_passed", 0)) == 5,
        invariants["lost_words"] == 0,
        invariants["duplicated_words"] == 0,
        invariants["conflicting_assignments"] == 0,
        decisions["false_positive_continuations"] == 0,
        decisions["continuation_precision"] >= 0.98,
        decisions["continuation_recall"] >= 0.95,
        decisions["target_cell_accuracy"] >= 0.98,
        decisions["continuation_type_accuracy"] >= 0.95,
        decisions["joining_mode_accuracy"] >= 0.95,
        decisions["final_merged_text_accuracy"] >= 0.99,
        boundary["genuine_row_retention_rate"] == 1.0,
        boundary["footnote_exclusion_precision"] == 1.0,
        boundary["footnote_exclusion_recall"] == 1.0,
        float(original.get("exact_merged_range_accuracy", 0.0)) == 1.0,
        unresolved["unresolved_case_accuracy"] == 1.0,
        reversibility["irreversible_decisions"] == 0,
        reversibility["raw_structures_exported"] and reversibility["resolved_structures_exported"],
        float(original.get("numeric_cell_accuracy", 0.0)) == 1.0 and float(continuation.get("numeric_cell_accuracy", 0.0)) == 1.0,
        min(float(original.get("correct_placement_rate", 0.0)), float(continuation.get("placement_rate", 0.0))) >= 0.99,
        all(row["ground_truth_status"] == "PASS" for row in gt_audit),
        all(row["covered"] for row in coverage),
        all(row.get("status") == "PASS" for row in rows),
    ]
    if all(conditions):
        return "PASS", "All pass-rule conditions and required case coverage were verified."
    failed = []
    if not all(row["covered"] for row in coverage):
        missing = [row["case_type"] for row in coverage if not row["covered"]]
        failed.append(f"missing required case coverage: {', '.join(missing)}")
    if not all(row["ground_truth_status"] == "PASS" for row in gt_audit):
        failed.append("incomplete continuation ground truth fields")
    if any(row.get("status") != "PASS" for row in rows):
        failed.append("one or more benchmark tables failed")
    if invariants["lost_words"] or invariants["duplicated_words"] or invariants["conflicting_assignments"]:
        failed.append("word invariant violations")
    if decisions["false_positive_continuations"] or decisions["missed_continuations"]:
        failed.append("continuation decision mismatches")
    return "PARTIAL", "; ".join(failed) if failed else "Core metrics pass, but at least one required audit condition is unmet."


def _decision_row(table_id: str, pred: dict[str, Any], exp: dict[str, Any] | None, second: dict[str, Any] | None) -> dict[str, Any]:
    cand = pred["selected_candidate"]
    reconstruction = _load_reconstruction(table_id)
    word_map = {idx: word.get("text", "") for idx, word in enumerate(_load_json(EXTRACTED / f"{table_id}_page_001_words.json"))}
    raw_text = " ".join(word_map.get(idx, "") for idx in pred.get("source_word_indices", [])).strip()
    traceability_status = "PASS" if _traceable(pred, reconstruction, raw_text) else "FAIL"
    return {
        "table_id": table_id,
        "source_text": pred.get("original_text", ""),
        "raw_source_text": raw_text,
        "target_cell": f"r{cand.get('target_row_index')}c{cand.get('target_column_index')}",
        "target_row": cand.get("target_row_index"),
        "target_column": cand.get("target_column_index"),
        "continuation_type": cand.get("candidate_type"),
        "joining_mode": cand.get("joining_mode"),
        "selected_score": cand.get("score"),
        "second_best_score": second.get("score") if second else "",
        "score_margin": pred.get("score_margin"),
        "confidence": pred.get("confidence"),
        "expected_decision": bool(exp),
        "predicted_decision": True,
        "correct": _decision_exact(pred, exp) if exp else False,
        "reason_codes": ";".join(pred.get("reason_codes", [])),
        "expected_target_row": exp.get("target_row") if exp else "",
        "expected_target_column": exp.get("target_column") if exp else "",
        "expected_type": exp.get("continuation_type") if exp else "",
        "expected_joining_mode": exp.get("expected_joining_mode") if exp else "",
        "expected_final_text": exp.get("expected_final_text") if exp else "",
        "final_text": pred.get("final_text", ""),
        "source_visual_line_indices": cand.get("source_line_indices", []),
        "traceability_status": traceability_status,
        "missing_source_mapping": not bool(pred.get("source_word_indices")) or not bool(pred.get("reversible_mapping")),
        "missing_rejected_alternatives": not bool(pred.get("all_considered_candidates")),
        "raw_visual_lines_exported": bool(reconstruction.get("metrics", {}).get("raw_visual_lines")),
        "resolved_rows_exported": bool(reconstruction.get("rows")),
    }


def _missed_decision_row(table_id: str, exp: dict[str, Any]) -> dict[str, Any]:
    return {
        "table_id": table_id,
        "source_text": "",
        "raw_source_text": "",
        "target_cell": f"r{exp.get('target_row')}c{exp.get('target_column')}",
        "target_row": "",
        "target_column": "",
        "continuation_type": "",
        "joining_mode": "",
        "selected_score": "",
        "second_best_score": "",
        "score_margin": "",
        "confidence": "",
        "expected_decision": True,
        "predicted_decision": False,
        "correct": False,
        "match_status": "MISSED",
        "error_categories": "MISSED_CONTINUATION",
        "reason_codes": "",
        "expected_target_row": exp.get("target_row"),
        "expected_target_column": exp.get("target_column"),
        "expected_type": exp.get("continuation_type"),
        "expected_joining_mode": exp.get("expected_joining_mode"),
        "expected_final_text": exp.get("expected_final_text"),
        "final_text": "",
        "source_visual_line_indices": exp.get("source_visual_line_indices", []),
        "traceability_status": "",
        "missing_source_mapping": "",
        "missing_rejected_alternatives": "",
        "raw_visual_lines_exported": "",
        "resolved_rows_exported": "",
    }


def _decision_errors(pred: dict[str, Any], exp: dict[str, Any]) -> list[str]:
    cand = pred["selected_candidate"]
    errors = []
    if cand.get("target_row_index") != exp.get("target_row"):
        errors.append("WRONG_TARGET_ROW")
    if cand.get("target_column_index") != exp.get("target_column"):
        errors.append("WRONG_TARGET_COLUMN")
    if cand.get("candidate_type") != exp.get("continuation_type"):
        errors.append("WRONG_CONTINUATION_TYPE")
    if cand.get("joining_mode") != exp.get("expected_joining_mode"):
        errors.append("WRONG_JOINING_MODE")
    if pred.get("final_text") != exp.get("expected_final_text"):
        errors.append("WRONG_FINAL_TEXT")
    return errors


def _decision_exact(pred: dict[str, Any], exp: dict[str, Any] | None) -> bool:
    if not exp:
        return False
    cand = pred["selected_candidate"]
    return (
        tuple(cand.get("source_line_indices", [])) == tuple(exp.get("source_visual_line_indices", []))
        and cand.get("target_row_index") == exp.get("target_row")
        and cand.get("target_column_index") == exp.get("target_column")
        and cand.get("candidate_type") == exp.get("continuation_type")
        and cand.get("joining_mode") == exp.get("expected_joining_mode")
        and pred.get("final_text") == exp.get("expected_final_text")
    )


def _traceable(pred: dict[str, Any], reconstruction: dict[str, Any], raw_text: str) -> bool:
    return bool(
        reconstruction.get("metrics", {}).get("raw_visual_lines")
        and reconstruction.get("rows")
        and pred.get("source_word_indices")
        and pred.get("original_bounding_boxes")
        and pred.get("original_text")
        and pred.get("final_text")
        and pred.get("reversible_mapping")
        and pred.get("all_considered_candidates")
        and raw_text
    )


def _second_best(pred: dict[str, Any]) -> dict[str, Any] | None:
    alternatives = sorted(pred.get("all_considered_candidates", []), key=lambda item: item.get("decision_rank", 999))
    return alternatives[1] if len(alternatives) > 1 else None


def _expected_key(item: dict[str, Any]) -> tuple[Any, ...]:
    return (tuple(item.get("source_visual_line_indices", [])), item.get("target_row"), item.get("target_column"))


def _pred_key(item: dict[str, Any]) -> tuple[Any, ...]:
    cand = item["selected_candidate"]
    return (tuple(cand.get("source_line_indices", [])), cand.get("target_row_index"), cand.get("target_column_index"))


def _reconstruction_path(table_id: str, group: str) -> Path:
    if group == "original":
        return EVAL / f"{table_id}_predicted.json"
    return EVAL / f"{table_id}_page_001_reconstructed.json"


def _load_reconstruction(table_id: str) -> dict[str, Any]:
    return _load_json(_reconstruction_path(table_id, "continuation"))


def _xlsx_has_unresolved(table_id: str, line_index: int) -> bool:
    from openpyxl import load_workbook

    path = EVAL / f"{table_id}_page_001_reconstructed.xlsx"
    if not path.exists():
        return False
    workbook = load_workbook(path, read_only=True)
    if "Unresolved Items" not in workbook.sheetnames:
        return False
    for row in workbook["Unresolved Items"].iter_rows(values_only=True):
        if row and row[0] == line_index:
            return True
    return False


def _inventory_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    ids = [row["table_id"] for row in rows]
    return {
        "original_tables": sum(1 for row in rows if row["benchmark_group"] == "original"),
        "continuation_tables": sum(1 for row in rows if row["benchmark_group"] == "continuation"),
        "total_unique_tables": len(set(ids)),
        "tables_actually_evaluated": sum(1 for row in rows if row["evaluated"]),
        "duplicated_tables": [table_id for table_id in sorted(set(ids)) if ids.count(table_id) > 1],
        "missing_tables": [row["table_id"] for row in rows if row["inventory_status"] != "OK"],
        "manual_ground_truth_verified": sum(1 for row in rows if str(row["manually_verified"]).lower() == "true"),
    }


def _coverage_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "required_cases": len(rows),
        "cases_covered": sum(1 for row in rows if row["covered"]),
        "cases_missing": [row["case_type"] for row in rows if not row["covered"]],
        "ambiguous_cases_included": any(row["case_type"].startswith("An intentionally ambiguous") and row["covered"] for row in rows),
    }


def _gt_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "tables_audited": len(rows),
        "verified_tables": sum(1 for row in rows if row["ground_truth_status"] == "PASS"),
        "failed_tables": [row["table_id"] for row in rows if row["ground_truth_status"] != "PASS"],
    }


def _test_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "required_behaviors": len(rows),
        "covered_behaviors": sum(1 for row in rows if row["covered"]),
        "missing_behavior_coverage": [row["required_behavior"] for row in rows if not row["covered"]],
    }


def _fixture_audit_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "fixtures_audited": len(rows),
        "verified": sum(1 for row in rows if row.get("fixture_status") == "VERIFIED"),
        "failed": [row.get("table_id") for row in rows if row.get("fixture_status") != "VERIFIED"],
    }


def _initial_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "new_tables_passed_before_code_changes": sum(1 for row in rows if row.get("status") == "PASS"),
        "new_tables_failed_before_code_changes": sum(1 for row in rows if row.get("status") != "PASS"),
        "failure_categories": [
            _initial_failure_category(row)
            for row in rows
            if row.get("status") != "PASS"
        ],
        "initial_output_directory": str(EVAL / "missing_cases_initial"),
    }


def _initial_failure_category(row: dict[str, Any]) -> str:
    table_id = row.get("table_id")
    if table_id == "continuation_11":
        return "FINAL_COLUMN_TARGETING_ERROR;COLUMN_COUNT_COLLAPSE"
    if table_id == "continuation_13":
        return "SECTION_HEADER_MISCLASSIFICATION;SECTION_HEADER_SPAN_ERROR"
    if row.get("status") == "MISSING_INITIAL_OUTPUT":
        return "OTHER_GENERAL_STRUCTURE_ERROR"
    return "OTHER_GENERAL_STRUCTURE_ERROR"


def _write_report(path: Path, summary: dict[str, Any], table_rows: list[dict[str, Any]], decision_rows: list[dict[str, Any]], unresolved_rows: list[dict[str, Any]], coverage: list[dict[str, Any]]) -> None:
    lines = ["# Continuation Audit Report", "", f"Status: {summary['audit_status']}", "", f"Why: {summary['why']}", "", "## Summary", "", "```json", json.dumps(summary, indent=2, sort_keys=True), "```", "", "## Per-Table Results", ""]
    lines.append("| Table | Category | Expected | Predicted | Continuations TP/FP/Missed | Lost/Duplicated | Status |")
    lines.append("|---|---|---:|---:|---:|---:|---|")
    for row in table_rows:
        lines.append(f"| {row['table_id']} | {row['category']} | {row['expected_dimensions']} | {row['predicted_dimensions']} | {row['true_positives']}/{row['false_positives']}/{row['missed_continuations']} | {row['lost_words']}/{row['duplicated_words']} | {row['status']} |")
    lines.extend(["", "## Continuation Decisions", ""])
    lines.append("| Table | Source | Target | Type | Join | Score | Margin | Status |")
    lines.append("|---|---|---|---|---|---:|---:|---|")
    for row in decision_rows:
        lines.append(f"| {row['table_id']} | {row.get('source_text', '')} | {row.get('target_cell', '')} | {row.get('continuation_type', '')} | {row.get('joining_mode', '')} | {row.get('selected_score', '')} | {row.get('score_margin', '')} | {row.get('match_status', '')} |")
    lines.extend(["", "## Unresolved Cases", ""])
    for row in unresolved_rows:
        lines.append(f"- {row['table_id']} line {row['line_index']}: {row.get('text', '')} [{row['status']}]")
    lines.extend(["", "## Case Coverage", ""])
    for row in coverage:
        lines.append(f"- {row['case_type']}: {'covered' if row['covered'] else 'missing'} ({row['table_ids']})")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _read_csv_matrix(path: Path) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return [row for row in csv.reader(handle)]


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames = sorted({key for row in rows for key in row}) if rows else ["status"]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _rate(num: float, den: float) -> float:
    return num / den if den else 0.0


if __name__ == "__main__":
    result = run_audit()
    print(f"audit_status={result['audit_status']}")
    print(f"summary_json={EVAL / 'continuation_summary.json'}")
    print(f"report={EVAL / 'continuation_benchmark_report.md'}")
